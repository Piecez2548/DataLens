import json

import pytest
from fastapi.testclient import TestClient

from app.analysis import analyze
from app.main import app


def test_scores_and_statistics():
    result = analyze(b"amount,group\n10,A\n20,B\n20,B\n,C\n", "demo.csv")
    assert result["scores"] == {
        "Completeness": 87.5,
        "Consistency": 100,
        "Uniqueness": 75,
        "Validity": 100,
        "Overall": 90.62,
    }
    assert result["columns"][0]["stats"]["median"] == 20
    assert result["duplicates"] == 1
    assert result["preview"][3]["amount"] is None


def test_mixed_types_and_validity_are_distinct():
    result = analyze(b"n,date\n1,2024-01-01\n2,2024-02-30\n3,2024-01-03\ninf,2024-01-04\noops,bad\n", "mixed.csv")
    assert result["scores"]["Consistency"] == 80
    assert result["scores"]["Validity"] == 75
    assert result["columns"][0]["stats"]["count"] == 3


@pytest.mark.parametrize("content", [b"", b"a\n", b"a,b\n1,2,3", b"a\n\xff", b"a\n\x00"])
def test_reject_bad_csv(content):
    with pytest.raises(ValueError):
        analyze(content, "bad.csv")


def test_all_missing_and_single_value():
    result = analyze(b"a,b\n,\n", "empty-values.csv")
    assert result["scores"]["Consistency"] is None
    assert result["scores"]["Validity"] is None
    assert result["scores"]["Overall"] == 50
    assert analyze(b"a\n3\n", "one.csv")["columns"][0]["stats"]["std"] == 0


def test_api_upload():
    client = TestClient(app)
    assert client.get("/api/health").status_code == 200
    response = client.post("/api/analyze", files={"file": ("demo.csv", b"x,y\n1,2\n3,4")})
    assert response.status_code == 200
    assert response.json()["scatter"] == [{"x": 1, "y": 2}, {"x": 3, "y": 4}]
    assert client.post("/api/analyze", files={"file": ("bad.txt", b"x\n1")}).status_code == 400
    assert client.post("/api/analyze", files={"file": ("bad.csv", b"x\n")}).status_code == 400


def test_headerless_csv_keeps_first_record_and_warns_executives():
    content = b"1,Marie,100.00\n2,Alex,250.00\n3,Sam,125.00\n"
    result = analyze(content, "customers.csv")
    assert result["rows"] == 3
    assert result["header"]["generated_names"] is True
    assert result["columns"][0]["name"] == "column_1"
    assert result["preview"][0]["column_2"] == "Marie"
    assert result["executive"]["status"] == "Review needed"
    assert result["executive"]["issues"][0]["title"] == "Column names are missing"


def test_header_mode_can_override_detection():
    content = b"id,name\n1,Ada\n2,Grace\n"
    result = analyze(content, "people.csv", "absent")
    assert result["rows"] == 3
    assert result["header"]["used"] is False
    assert result["preview"][0] == {"column_1": "id", "column_2": "name"}


def test_executive_summary_prioritizes_invalid_values():
    result = analyze(b"date\n2026-02-30\n2026-01-01\n", "dates.csv", "present")
    assert result["executive"]["status"] == "Action required"
    assert result["executive"]["high_priority_count"] == 1


def test_api_accepts_header_override():
    client = TestClient(app)
    response = client.post(
        "/api/analyze?header_mode=absent",
        files={"file": ("demo.csv", b"1,A\n2,B\n")},
    )
    assert response.status_code == 200
    assert response.json()["rows"] == 2
    assert response.json()["header"]["generated_names"] is True


def test_semantic_email_and_identifier_detection_improve_validity_scope():
    result = analyze(
        b"customer_id,email\n1,ada@example.com\n2,invalid\n3,grace@example.com\n",
        "contacts.csv",
    )
    assert result["columns"][0]["type"] == "identifier"
    assert result["columns"][1]["type"] == "email"
    assert result["columns"][1]["invalid"] == 1
    assert result["scores"]["Validity"] == 66.67
    assert result["executive"]["status"] == "Action required"


def test_iqr_outliers_and_strongest_correlation_are_reported():
    result = analyze(
        b"amount,revenue\n1,10\n2,20\n3,30\n4,40\n100,1000\n",
        "metrics.csv",
    )
    assert result["columns"][0]["outliers"] == 1
    assert result["correlations"][0]["coefficient"] > 0.99
    assert result["scatter_axes"] == ["amount", "revenue"]
    assert result["executive"]["signals"]


def test_schema_configuration_names_columns_without_dropping_first_row():
    result = analyze(
        b"1,ada@example.com\n2,grace@example.com\n",
        "contacts.csv",
        "absent",
        ["customer_id", "email"],
        {"customer_id": "identifier", "email": "email"},
    )
    assert result["rows"] == 2
    assert result["header"]["configured_names"] is True
    assert result["header"]["generated_names"] is False
    assert [column["type"] for column in result["columns"]] == ["identifier", "email"]


def test_api_accepts_schema_configuration():
    client = TestClient(app)
    schema = json.dumps({"column_names": ["customer_id", "email"], "type_overrides": {"email": "email"}})
    response = client.post(
        "/api/analyze?header_mode=absent",
        files={"file": ("contacts.csv", b"1,ada@example.com\n2,bad\n"), "schema": (None, schema)},
    )
    assert response.status_code == 200
    assert response.json()["columns"][1]["invalid"] == 1


def test_provenance_fingerprints_the_exact_input_without_generating_values():
    content = b"id,value\n1,10\n2,20\n"
    result = analyze(content, "evidence.csv")
    assert result["provenance"] == {
        "source": "uploaded_file",
        "sha256": "159f8bae5fa563fb61b540de391850552f9fe1d188273b1ca9ce182e4cbf4a26",
        "input_bytes": len(content),
        "file_rows": 3,
        "blank_rows_skipped": 0,
        "analyzed_rows": 2,
        "preview_rows": 2,
        "method_version": "0.7.1",
        "source_format": "csv",
        "sheet": None,
        "sheets": None,
        "header_row": 1,
        "header_row_suggestion": None,
        "date_order": "dmy",
        "date_order_ambiguous": False,
        "two_digit_years": 0,
        "year_era": None,
        "ignored_warnings": [],
        "encoding": "utf-8",
        "delimiter": "comma",
        "missing_markers": [],
        "calculation_mode": "deterministic",
        "data_values_generated": False,
    }


def test_embedded_table_delimiters_require_structural_review():
    result = analyze(
        b'ID,| Category | Product |\n1,"| Drinks | Water |"\n2,"| Food | Rice |"\n',
        "products.csv",
    )
    assert result["scores"]["Overall"] == 100
    assert result["executive"]["status"] == "Review needed"
    assert any("embedded" in issue["title"].lower() for issue in result["executive"]["issues"])


def test_iqr_flags_minority_values_when_middle_half_is_constant():
    result = analyze(b"volume\n180\n180\n180\n180\n200\n", "volume.csv")
    assert result["columns"][0]["stats"]["q1"] == 180
    assert result["columns"][0]["stats"]["q3"] == 180
    assert result["columns"][0]["outliers"] == 1


def test_business_rules_block_readiness():
    result = analyze(
        b"region,amount\nTH,10\nUS,200\n,50\n",
        "rules.csv",
        business_rules={
            "region": {"required": True, "allowed_values": ["TH"]},
            "amount": {"min": 20, "max": 100},
        },
    )
    assert result["business_rules"]["total_violations"] == 4
    assert result["business_rules"]["passed"] is False
    assert result["executive"]["status"] == "Action required"


def test_analysis_needs_no_account_and_nothing_is_cached():
    response = TestClient(app).post("/api/analyze", files={"file": ("demo.csv", b"x\n1\n")})
    body = response.json()
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store, max-age=0"
    assert "governance" not in body
    assert "audit_event" not in body


def test_health_reports_no_authentication_requirement():
    assert TestClient(app).get("/api/health").json() == {"status": "ok", "version": app.version}


def test_hosted_analysis_requires_the_non_personal_data_attestation(monkeypatch):
    monkeypatch.setenv("VERCEL", "1")
    client = TestClient(app)
    refused = client.post("/api/analyze", files={"file": ("demo.csv", b"x\n1\n")})
    assert refused.status_code == 400
    schema = json.dumps({"authorized_to_process": True})
    accepted = client.post("/api/analyze", files={"file": ("demo.csv", b"x\n1\n"), "schema": (None, schema)})
    assert accepted.status_code == 200
    invalid = client.post(
        "/api/analyze",
        files={"file": ("demo.csv", b"x\n1\n"), "schema": (None, json.dumps({"authorized_to_process": "yes"}))},
    )
    assert invalid.status_code == 400


def test_approval_endpoints_no_longer_exist():
    assert TestClient(app).post("/api/audit/review", json={}).status_code in {404, 405}


def test_numeric_business_rule_rejects_a_text_column():
    with pytest.raises(ValueError, match="require a numeric column"):
        analyze(
            b"region\nThailand\nJapan\n",
            "regions.csv",
            business_rules={"region": {"min": 0}},
        )


def test_auto_header_does_not_drop_an_ambiguous_text_row():
    result = analyze(b"Al,NY\nCharlotte,California\nBenjamin,Washington\n", "places.csv")
    assert result["header"]["used"] is False
    assert result["rows"] == 3
    assert result["preview"][0] == {"column_1": "Al", "column_2": "NY"}


@pytest.mark.parametrize(
    ("content", "expected_type"),
    [
        (b"amount\n10\n20\noops\n", "numeric"),
        (b"date\n2024-01-01\n2024-01-02\nbad\n", "datetime"),
        (b"flag\ntrue\nfalse\nmaybe\n", "boolean"),
    ],
)
def test_dominant_typed_values_expose_dirty_minority(content, expected_type):
    result = analyze(content, "dirty.csv", "present")
    assert result["columns"][0]["type"] == expected_type
    assert result["columns"][0]["mismatches"] == 1
    assert result["scores"]["Overall"] < 100
    assert result["executive"]["status"] == "Review needed"


def test_sequential_business_measures_remain_numeric():
    result = analyze(b"age,sales\n20,100\n21,101\n22,102\n23,103\n", "measures.csv")
    assert [column["type"] for column in result["columns"]] == ["numeric", "numeric"]
    assert result["columns"][0]["stats"]["mean"] == 21.5
    assert result["correlations"]


@pytest.mark.parametrize(("delimiter", "label"), [(b";", "semicolon"), (b"\t", "tab"), (b"|", "pipe")])
def test_alternate_delimiters_are_detected(delimiter, label):
    content = delimiter.join([b"name", b"amount"]) + b"\n" + delimiter.join([b"A", b"10"]) + b"\n"
    result = analyze(content, "alt.csv")
    assert result["column_count"] == 2
    assert result["provenance"]["delimiter"] == label


def test_thai_windows_874_and_utf16_files_are_decoded():
    text = "ชื่อ,จำนวน\nสมชาย,10\nสมหญิง,20\n"
    legacy = analyze(text.encode("cp874"), "th.csv")
    assert legacy["provenance"]["encoding"] == "cp874"
    assert legacy["columns"][0]["name"] == "ชื่อ"
    utf16 = analyze(text.encode("utf-16"), "th16.csv")
    assert utf16["provenance"]["encoding"] == "utf-16"
    assert utf16["preview"][0]["ชื่อ"] == "สมชาย"
    assert analyze(text.encode("utf-8-sig"), "bom.csv")["provenance"]["encoding"] == "utf-8"


def test_undecodable_bytes_are_rejected():
    with pytest.raises(ValueError, match="encoded CSV"):
        analyze(b"a,b\n\xd8\xff,1\n\x81\x8d,2\n", "bad.csv")


def test_missing_issue_severity_uses_all_cells_as_denominator():
    headers = ",".join(f"c{i}" for i in range(100))
    values = ",".join([""] + ["x"] * 99)
    result = analyze(f"{headers}\n{values}\n".encode(), "wide.csv", "present")
    missing_issue = next(issue for issue in result["executive"]["issues"] if "missing" in issue["title"])
    assert result["scores"]["Completeness"] == 99
    assert missing_issue["severity"] == "medium"


def test_business_rule_findings_are_never_truncated():
    result = analyze(
        b'id,date,email,amount,embedded\n1,2024-02-30,bad,1,"| a | b |"\n1,2024-02-30,bad,,"| a | b |"\n100,2024-01-01,a@example.com,1000,"| a | b |"\n',
        "many-issues.csv",
        business_rules={"amount": {"required": True, "max": 100}},
    )
    assert len(result["executive"]["issues"]) == result["executive"]["issue_count"]
    assert any("business-rule" in issue["title"] for issue in result["executive"]["issues"])


def test_empty_or_disabled_rules_are_not_reported_as_configured():
    assert analyze(b"amount\n10\n", "rules.csv", business_rules={"amount": {}})["business_rules"]["configured"] is False
    assert analyze(b"amount\n10\n", "rules.csv", business_rules={"amount": {"required": False}})["business_rules"]["configured"] is False


def test_numeric_rule_rejects_boolean_and_inverted_range():
    with pytest.raises(ValueError, match="finite numbers"):
        analyze(b"amount\n10\n", "rules.csv", business_rules={"amount": {"min": True}})
    with pytest.raises(ValueError, match="no greater than maximum"):
        analyze(b"amount\n10\n", "rules.csv", business_rules={"amount": {"min": 20, "max": 10}})


def test_blank_line_in_a_one_column_file_is_counted_as_a_missing_value():
    result = analyze(b"value\n1\n\n2\n\n\n", "blank-row.csv", "present")
    assert result["rows"] == 3
    assert result["missing"] == 1


def test_separator_lines_and_empty_rows_in_a_table_are_skipped_but_row_numbers_stay():
    result = analyze(b"id,v\n1,\n\n,\n2,\n3,9\n", "gaps.csv", "present")
    assert result["rows"] == 3
    assert result["provenance"]["blank_rows_skipped"] == 2
    rows = [e["row"] for e in result["issue_samples"]["kinds"]["missing"]["examples"]]
    assert rows == [2, 5]


def test_headers_are_repaired_instead_of_refusing_the_file():
    trailing = analyze(b"id,v,\n1,2,\n3,4,\n", "t.csv")
    assert [c["name"] for c in trailing["columns"]] == ["id", "v"]
    assert trailing["header"]["blank_columns_removed"] == 1
    renamed = analyze(b"id,total,total,\n1,2,3,x\n4,5,6,y\n", "r.csv")
    assert [c["name"] for c in renamed["columns"]] == ["id", "total", "total_2", "column_4"]
    assert [item["to"] for item in renamed["header"]["repaired"]] == ["total_2", "column_4"]


def test_lone_carriage_return_line_endings_are_read():
    assert analyze(b"id,v\r1,2\r3,4\r", "mac.csv")["rows"] == 2


def _workbook(rows, sheet_title="Data"):
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_title
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_missing_markers_count_literal_placeholders_as_missing():
    content = b"name,qty\nA,1\nB,N/A\nC,-\nD,null\n"
    assert analyze(content, "m.csv")["missing"] == 0
    result = analyze(content, "m.csv", missing_markers=["n/a", " - ", "NULL"])
    assert result["missing"] == 3
    assert result["provenance"]["missing_markers"] == ["-", "n/a", "null"]
    assert result["columns"][1]["type"] == "numeric"


def test_missing_markers_are_validated():
    with pytest.raises(ValueError, match="Missing markers"):
        analyze(b"a\n1\n", "m.csv", missing_markers=["x"] * 21)
    with pytest.raises(ValueError, match="Missing markers"):
        analyze(b"a\n1\n", "m.csv", missing_markers=[""])


def test_uneven_rows_are_located():
    with pytest.raises(ValueError, match=r"expected 2, but row 3 has 3, row 4 has 1"):
        analyze(b"a,b\n1,2\n3,4,5\n6\n", "uneven.csv")


def test_xlsx_first_sheet_is_read_with_dates_and_numbers():
    import datetime

    content = _workbook(
        [
            ["order_date", "amount", "paid"],
            [datetime.date(2024, 1, 5), 10.0, True],
            [datetime.date(2024, 3, 9), 12.5, False],
        ],
        "Sales",
    )
    result = analyze(content, "book.xlsx")
    assert result["rows"] == 2
    assert [c["type"] for c in result["columns"]] == ["datetime", "numeric", "boolean"]
    assert result["provenance"]["source_format"] == "xlsx"
    assert result["provenance"]["sheet"] == "Sales"
    assert result["provenance"]["delimiter"] is None
    assert len(result["provenance"]["sha256"]) == 64


def test_api_accepts_xlsx_json_and_sheet_choice():
    import json

    client = TestClient(app)
    book = _workbook([["a", "b"], [1, 2], [3, 4]], "Only")
    ok = client.post("/api/analyze", files={"file": ("b.xlsx", book)}, data={"schema": json.dumps({"sheet": "Only"})})
    assert ok.status_code == 200 and ok.json()["provenance"]["sheet"] == "Only"
    bad = client.post("/api/analyze", files={"file": ("b.xlsx", book)}, data={"schema": json.dumps({"sheet": "X"})})
    assert bad.status_code == 400
    js = client.post("/api/analyze", files={"file": ("r.json", b'[{"a": 1}, {"a": 2}]')})
    assert js.status_code == 200 and js.json()["rows"] == 2


def test_xlsx_sheet_can_be_selected():
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    workbook.active.append(["a"])
    workbook.active.append([1])
    second = workbook.create_sheet("Detail")
    second.append(["x", "y"])
    second.append([1, 2])
    second.append([3, 4])
    buffer = io.BytesIO()
    workbook.save(buffer)
    first = analyze(buffer.getvalue(), "b.xlsx")
    assert first["provenance"]["sheets"] == ["Summary", "Detail"]
    assert first["column_count"] == 1
    chosen = analyze(buffer.getvalue(), "b.xlsx", sheet="Detail")
    assert (chosen["provenance"]["sheet"], chosen["rows"], chosen["column_count"]) == ("Detail", 2, 2)
    with pytest.raises(ValueError, match="does not exist"):
        analyze(buffer.getvalue(), "b.xlsx", sheet="Nope")


def test_json_records_are_flattened_with_the_union_of_keys():
    content = (
        b'{"data": [{"id": 1, "user": {"name": "A", "age": 30}, "tags": ["x"]},'
        b' {"id": 2, "user": {"name": "B"}, "extra": true}]}'
    )
    result = analyze(content, "records.json")
    names = [c["name"] for c in result["columns"]]
    assert names == ["id", "user.name", "user.age", "tags", "extra"]
    assert result["rows"] == 2
    assert result["provenance"]["source_format"] == "json"
    assert result["issue_samples"]["row_label"] == "record"
    assert result["columns"][2]["missing"] == 1


def test_json_lines_and_errors():
    assert analyze(b'{"a": 1}\n{"a": 2}\n', "lines.jsonl")["rows"] == 2
    with pytest.raises(ValueError, match="not valid JSON"):
        analyze(b"{broken", "bad.json")
    with pytest.raises(ValueError, match="non-empty array"):
        analyze(b"[]", "empty.json")


def test_issue_samples_locate_rows_with_file_numbering():
    content = b"id,amount,when\n1,10,2024-01-01\n2,,2024-02-30\n3,oops,2024-01-03\n3,oops,2024-01-03\n4,1000,2024-01-04\n5,20,2024-01-05\n6,30,2024-01-06\n"
    result = analyze(content, "loc.csv", "present", business_rules={"amount": {"min": 0, "max": 100}})
    kinds = result["issue_samples"]["kinds"]
    assert result["issue_samples"]["row_label"] == "row"
    assert [(e["row"], e["column"]) for e in kinds["missing"]["examples"]] == [(3, "amount")]
    assert [e["row"] for e in kinds["invalid"]["examples"]] == [3]
    assert [e["row"] for e in kinds["mismatch"]["examples"]] == [4, 5]
    assert [e["row"] for e in kinds["duplicate"]["examples"]] == [5]
    assert [e["row"] for e in kinds["rule"]["examples"]] == [6]
    assert kinds["rule"]["examples"][0]["reason"] == "Above the maximum 100"


def test_issue_samples_are_capped_but_totals_are_exact():
    content = b"v,w\n" + b"".join(b",1\n" for _ in range(3000))
    result = analyze(content, "many.csv", "present")
    entry = result["issue_samples"]["kinds"]["missing"]
    assert entry["total"] == 3000
    assert len(entry["examples"]) == 1000


def test_money_percent_and_accounting_numbers_are_numeric():
    content = '''item,price,share,delta
a,"1,234.50",10%,(500)
b,฿100,15.5%,-20
c,"$2,000",20%,30
d,300 บาท,25%,40
'''.encode()
    result = analyze(content, "money.csv", business_rules={"price": {"min": 0}})
    price, share, delta = result["columns"][1:]
    assert (price["type"], price["format"], price["mismatches"]) == ("numeric", "currency", 0)
    assert price["stats"]["max"] == 2000.0 and price["stats"]["min"] == 100.0
    assert (share["type"], share["format"]) == ("numeric", "percent")
    assert delta["stats"]["min"] == -500.0
    assert result["business_rules"]["passed"] is True


def test_malformed_money_is_not_numeric():
    result = analyze(b'v\n"1,2"\n"12,34,5"\n$$5\n--3\n', "bad.csv", "present")
    assert result["columns"][0]["type"] == "categorical"


def test_date_order_is_inferred_from_proof_and_flagged_when_ambiguous():
    assert analyze(b"d\n31/12/2024\n01/02/2024\n", "a.csv", "present")["provenance"]["date_order"] == "dmy"
    mdy = analyze(b"d\n12/31/2024\n01/02/2024\n", "b.csv", "present")
    assert mdy["provenance"]["date_order"] == "mdy"
    assert mdy["columns"][0]["invalid"] == 0
    ambiguous = analyze(b"d\n05/06/2024\n01/02/2024\n", "c.csv", "present")
    assert ambiguous["provenance"]["date_order_ambiguous"] is True
    forced = analyze(b"d\n05/06/2024\n01/02/2024\n", "c.csv", "present", date_order="mdy")
    assert forced["provenance"]["date_order"] == "mdy"
    assert forced["provenance"]["date_order_ambiguous"] is False
    assert forced["timeseries"]["first"] == "2024-01-02"
    with pytest.raises(ValueError, match="date_order"):
        analyze(b"d\n1\n", "c.csv", date_order="ymd")


RAGGED = b"a,b\n1,2,3\n4,5\n"


def test_header_row_skips_report_titles():
    content = b"Sales report\n\nregion,total\nN,1\nS,2\n"
    with pytest.raises(ValueError, match="set the header row"):
        analyze(RAGGED, "ragged.csv")
    automatic = analyze(content, "report.csv")
    assert automatic["provenance"]["header_row"] == 3
    assert automatic["rows"] == 2
    result = analyze(content, "report.csv", header_row=3)
    assert [c["name"] for c in result["columns"]] == ["region", "total"]
    assert result["rows"] == 2
    assert result["provenance"]["header_row"] == 3
    # row numbers still refer to the original file
    result = analyze(b"x\nt\nregion,total\nN,1\nS,\n".replace(b"x\nt\n", b"x,y\nt,u\n"), "r.csv", "present", header_row=3)
    assert result["issue_samples"]["kinds"]["missing"]["examples"][0]["row"] == 5
    with pytest.raises(ValueError, match="header row"):
        analyze(content, "report.csv", header_row=99)


def test_header_row_is_suggested_for_padded_workbooks():
    book = _workbook([["Quarterly report"], [], ["region", "total", "owner"], ["N", 1, "A"], ["S", 2, "B"]])
    result = analyze(book, "q.xlsx")
    assert result["provenance"]["header_row_suggestion"] == 3
    fixed = analyze(book, "q.xlsx", header_row=3)
    assert [c["name"] for c in fixed["columns"]] == ["region", "total", "owner"]


def test_business_key_duplicates_are_found_even_when_other_values_differ():
    content = b"po,qty,note\nA1,5,x\nA2,6,y\nA1,9,z\n,1,blank\n,2,blank\nA1,1,w\n"
    assert analyze(content, "k.csv", "present")["duplicates"] == 0
    result = analyze(content, "k.csv", "present", business_key=["po"])
    assert result["key_duplicates"] == {"columns": ["po"], "duplicate_rows": 2, "blank_key_rows": 2}
    examples = result["issue_samples"]["kinds"]["key_duplicate"]["examples"]
    assert [e["row"] for e in examples] == [4, 7]
    assert examples[0]["reason"] == "Repeats a key first seen at row 2"
    assert result["executive"]["status"] == "Action required"
    with pytest.raises(ValueError, match="business key"):
        analyze(content, "k.csv", "present", business_key=["nope"])


def test_group_summary_ranks_groups_by_measure():
    content = b"team,amount\nB,10\nA,5\nB,20\nA,\nC,1\n,7\n"
    result = analyze(content, "g.csv", "present", group_by={"by": "team", "measure": "amount"})
    summary = result["group_summary"]
    assert [row["group"] for row in summary["rows"]] == ["B", "(blank)", "A", "C"]
    assert summary["rows"][0] == {"group": "B", "rows": 2, "sum": 30.0, "mean": 15.0, "min": 10.0, "max": 20.0}
    assert summary["total_groups"] == 4 and summary["truncated"] is False
    counts = analyze(content, "g.csv", "present", group_by={"by": "team"})["group_summary"]
    assert counts["rows"][0]["group"] == "B" and counts["rows"][0]["sum"] is None
    with pytest.raises(ValueError, match="numeric"):
        analyze(content, "g.csv", "present", group_by={"by": "amount", "measure": "team"})


def test_corrupt_xlsx_is_rejected():
    with pytest.raises(ValueError, match="could not be read"):
        analyze(b"PK\x03\x04not really a workbook", "bad.xlsx")


def test_thai_national_id_is_validated_and_flagged_as_personal_data():
    # 1101700230708 has a correct check digit; 1101700230706 does not.
    result = analyze("เลขบัตรประชาชน,จังหวัด\n1101700230708,A\n1101700230706,B\n".encode(), "ids.csv")
    column = result["columns"][0]
    assert column["pattern"] == "thai_national_id"
    assert column["type"] == "identifier"
    assert column["invalid"] == 1
    assert result["executive"]["status"] == "Action required"
    assert any("personal data" in issue["title"] for issue in result["executive"]["issues"])


def test_epoch_milliseconds_are_not_mistaken_for_national_ids():
    result = analyze(b"ts\n1700000000000\n1700000100000\n1700000200000\n", "ts.csv")
    assert result["columns"][0]["pattern"] is None


def test_thai_phone_numbers_are_flagged_as_personal_data():
    result = analyze(b"phone,v\n0812345678,1\n029876543,2\n", "phones.csv")
    assert result["columns"][0]["pattern"] == "thai_phone"
    assert result["executive"]["high_priority_count"] == 1


def test_timeseries_counts_periods_and_gaps():
    content = b"d\n2024-01-05\n2024-01-20\n2024-06-02\n2024-06-03\n"
    series = analyze(content, "ts.csv")["timeseries"]
    assert series["granularity"] == "month"
    assert [point["count"] for point in series["points"]] == [2, 0, 0, 0, 0, 2]
    assert series["gaps"] == 4
    assert (series["first"], series["last"]) == ("2024-01-05", "2024-06-03")
    assert analyze(b"x\n1\n", "n.csv")["timeseries"] is None


def test_day_first_slash_and_buddhist_era_dates_are_validated():
    result = analyze(
        b"date\n31/12/2024\n01/02/2567\n2024/03/15\n2024-05-06T10:30:00\n31/02/2024\n",
        "dates.csv",
    )
    column = result["columns"][0]
    assert column["type"] == "datetime"
    assert column["mismatches"] == 0
    assert column["invalid"] == 1


def test_executive_summary_is_provided_in_thai_without_changing_the_status():
    result = analyze(b"a,b\n1,\n2,3\n", "th.csv", "present")
    executive = result["executive"]
    assert executive["status"] == "Action required" or executive["status"] == "Review needed"
    thai = executive["thai"]
    assert len(thai["issues"]) == len(executive["issues"])
    assert "ช่อง" in thai["issues"][0]["title"]
    assert thai["scope_note"] != executive["scope_note"]


def test_european_numbers_are_numeric_but_english_thousands_are_not_misread():
    content = 'v\n"1.234,50"\n"12,5"\n"1 234,5"\n"€2.000,00"\n'.encode()
    column = analyze(content, "eu.csv", "present")["columns"][0]
    assert column["type"] == "numeric"
    assert (column["stats"]["min"], column["stats"]["max"]) == (12.5, 2000.0)
    english = analyze(b'v\n"1,234"\n"2,500"\n', "en.csv", "present")["columns"][0]
    assert english["stats"]["max"] == 2500.0


def test_named_month_dates_in_english_and_thai_are_dates():
    content = "d\n5 Jan 2024\n05-Feb-2024\n\"March 7, 2024\"\n9 มกราคม 2567\n10 ก.พ. 2567\n31 Feb 2024\n".encode()
    column = analyze(content, "named.csv", "present")["columns"][0]
    assert column["type"] == "datetime"
    assert column["mismatches"] == 0
    assert column["invalid"] == 1


def test_two_digit_years_are_reported_and_resolved_only_when_the_era_is_chosen():
    csv_bytes = b"id,date\n1,05/06/67\n2,25/12/67\n3,01/01/68\n"
    plain = analyze(csv_bytes, "d.csv")
    assert plain["provenance"]["two_digit_years"] == 3
    assert plain["columns"][1]["type"] != "datetime"
    buddhist = analyze(csv_bytes, "d.csv", year_era="be")
    assert buddhist["columns"][1]["type"] == "datetime"
    assert buddhist["columns"][1]["invalid"] == 0
    assert buddhist["timeseries"]["first"] == "2024-06-05"
    christian = analyze(csv_bytes, "d.csv", year_era="ce")
    assert christian["timeseries"]["first"] == "2067-06-05"


def test_year_era_must_be_known():
    import pytest

    with pytest.raises(ValueError):
        analyze(b"a,b\n1,2\n", "d.csv", year_era="xx")


def test_excel_date_serials_in_a_plain_number_column_are_flagged():
    lines = ["order,due_date,amount"] + [f"{i},{45366 + i},{i * 10}" for i in range(1, 6)]
    result = analyze(chr(10).join(lines).encode(), "serial.csv")
    titles = [item["title"] for item in result["executive"]["issues"]]
    assert "Numbers that look like Excel dates" in titles
    plain = analyze(chr(10).join(["order,amount"] + [f"{i},{45366 + i}" for i in range(1, 6)]).encode(), "plain.csv")
    assert "Numbers that look like Excel dates" not in [item["title"] for item in plain["executive"]["issues"]]


def test_problem_rows_return_the_original_rows_behind_the_issues():
    result = analyze(b"id,team,amount\n1,North,10\n2,,20\n3,South,\n3,South,\n", "a.csv")
    rows = result["problem_rows"]
    assert rows["total"] == 3
    assert rows["columns"] == ["id", "team", "amount"]
    assert [item["row"] for item in rows["rows"]] == [3, 4, 5]
    assert rows["rows"][0]["values"] == ["2", "", "20"]
    assert "missing:team" in rows["rows"][0]["issues"]
    clean = analyze(b"id,amount\n1,10\n2,20\n", "b.csv")
    assert clean["problem_rows"]["total"] == 0


def test_gzip_upload_is_unpacked_and_analyzed_under_its_original_name():
    import gzip

    body = gzip.compress(b"id,amount\n1,10\n2,20\n")
    response = TestClient(app).post("/api/analyze", files={"file": ("data.csv.gz", body)})
    assert response.status_code == 200
    assert response.json()["filename"] == "data.csv"
    assert response.json()["rows"] == 2


def test_gzip_upload_rejects_bad_data_and_decompression_bombs(monkeypatch):
    import gzip

    client = TestClient(app)
    assert client.post("/api/analyze", files={"file": ("x.csv.gz", b"not gzip")}).status_code == 400
    assert client.post("/api/analyze", files={"file": ("x.xlsx.gz", gzip.compress(b"x"))}).status_code == 400
    monkeypatch.setattr("app.main.MAX_BYTES", 1000)
    bomb = gzip.compress(b"a,b\n" + b"1,2\n" * 5000)
    assert client.post("/api/analyze", files={"file": ("big.csv.gz", bomb)}).status_code == 413


def test_files_with_too_many_cells_are_refused_with_a_clear_message(monkeypatch):
    monkeypatch.setattr("app.analysis.MAX_CELLS", 10)
    with pytest.raises(ValueError, match="too many cells"):
        analyze(b"a,b,c\n1,2,3\n4,5,6\n7,8,9\n10,11,12\n", "wide.csv")


def test_redact_removes_the_chosen_columns_and_nothing_else():
    client = TestClient(app)
    csv_bytes = b"id,phone,amount\n1,0812345678,10\n2,0898765432,20\n"
    schema = json.dumps({"drop_columns": ["phone"]})
    response = client.post(
        "/api/redact?header_mode=present",
        files={"file": ("p.csv", csv_bytes), "schema": (None, schema)},
    )
    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    assert text.splitlines() == ["id,amount", "1,10", "2,20"]
    assert "0812345678" not in text
    missing = client.post(
        "/api/redact",
        files={"file": ("p.csv", csv_bytes), "schema": (None, json.dumps({"drop_columns": ["nope"]}))},
    )
    assert missing.status_code == 400
    everything = client.post(
        "/api/redact?header_mode=present",
        files={"file": ("p.csv", b"a\n1\n"), "schema": (None, json.dumps({"drop_columns": ["a"]}))},
    )
    assert everything.status_code == 400


def test_type_conflicts_name_the_offending_values():
    result = analyze(b"id,amount\n1,10\n2,N/A\n3,20\n4,oops\n5,30\n", "m.csv")
    detail = next(i["detail"] for i in result["executive"]["issues"] if "conflict" in i["title"])
    assert "amount" in detail and "'N/A'" in detail


def test_context_warnings_flag_future_dates_and_negative_amounts():
    result = analyze(b"id,order_date,amount\n1,2024-01-05,100\n2,2099-01-01,-5\n3,2024-02-01,50\n", "c.csv")
    titles = " | ".join(i["title"] for i in result["executive"]["issues"])
    assert "in order_date are in the future" in titles
    assert "negative values in amount" in titles
    quiet = analyze(b"id,note_date,balance\n1,2099-01-01,-5\n2,2024-02-01,-6\n", "q.csv")
    assert not any("future" in i["title"] or "negative" in i["title"] for i in quiet["executive"]["issues"])


def test_spelling_variants_are_flagged_by_case_and_spacing_only():
    result = analyze(b"id,city\n1,Bangkok\n2,bangkok \n3,Bangkok\n4,Chiang Mai\n5,Phuket\n", "v.csv")
    variant = result["issue_samples"]["kinds"]["variant"]
    assert variant["total"] == 1
    assert variant["examples"][0]["row"] == 3
    assert any("spelled differently" in i["title"] for i in result["executive"]["issues"])
    same = analyze(b"id,city\n1,Bangkok\n2,Chiang Mai\n3,Bangkok\n", "s.csv")
    assert "variant" not in same["issue_samples"]["kinds"]


def test_ignored_warnings_leave_the_verdict_and_findings_untouched_by_them():
    content = b"id,amount\n1,10\n2,-5\n3,20\n"
    base = analyze(content, "r.csv")
    code = next(i["code"] for i in base["executive"]["issues"] if "negative" in i["title"])
    assert code == "negative_amount:amount"
    assert base["executive"]["status"] == "Review needed"
    quiet = analyze(content, "r.csv", ignored_warnings=[code])
    assert quiet["executive"]["status"] == "Ready for exploration"
    assert quiet["provenance"]["ignored_warnings"] == [code]
    with pytest.raises(ValueError, match="Ignored warnings"):
        analyze(content, "r.csv", ignored_warnings="all")


def test_core_findings_cannot_be_ignored():
    result = analyze(b"id,team\n1,a\n2,\n", "m.csv", ignored_warnings=["missing", "everything"])
    assert any("missing" in item["title"] for item in result["executive"]["issues"])


def _xlsx(rows, hide_column=None, hide_row=None, hidden_sheet=False):
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    if hide_column:
        sheet.column_dimensions[hide_column].hidden = True
    if hide_row:
        sheet.row_dimensions[hide_row].hidden = True
    if hidden_sheet:
        workbook.create_sheet("Notes").sheet_state = "hidden"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_hidden_columns_rows_and_sheets_are_reported_but_still_checked():
    rows = [["id", "secret", "amount"]] + [[i, f"s{i}", i * 10] for i in range(5)]
    result = analyze(_xlsx(rows, hide_column="B", hide_row=3, hidden_sheet=True), "h.xlsx")
    detail = next(i["detail"] for i in result["executive"]["issues"] if "Hidden content" in i["title"])
    assert "secret" in detail and "Notes" in detail and "1 hidden row" in detail
    assert result["rows"] == 5
    visible = analyze(_xlsx(rows), "v.xlsx")
    assert not any("Hidden" in i["title"] for i in visible["executive"]["issues"])


def test_numbers_stored_as_text_in_excel_are_flagged_but_codes_with_leading_zeros_are_not():
    rows = [["id", "qty", "code"]] + [[i, str(i * 3), f"00{i}"] for i in range(1, 8)]
    result = analyze(_xlsx(rows), "t.xlsx")
    titles = [i["title"] for i in result["executive"]["issues"]]
    assert "Numbers stored as text in qty" in titles
    assert next(c for c in result["columns"] if c["name"] == "qty")["type"] == "numeric"


def test_a_totals_row_that_equals_the_sum_above_is_flagged():
    flagged = analyze(b"id,team,amount\n1,a,10\n2,b,20\n3,a,30\nTotal,,60\n", "t.csv")
    assert any("looks like a totals row" in i["title"] for i in flagged["executive"]["issues"])
    assert flagged["issue_samples"]["kinds"]["totals"]["examples"][0]["row"] == 5
    not_sum = analyze(b"id,team,amount\n1,a,10\n2,b,20\n3,a,30\nTotal,,99\n", "t.csv")
    assert "totals" not in not_sum["issue_samples"]["kinds"]
    thai = analyze("id,team,amount\n1,a,10\n2,b,20\nรวม,,30\n".encode(), "t.csv")
    assert "totals" in thai["issue_samples"]["kinds"]
