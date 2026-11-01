import datetime
import io

import openpyxl
from openpyxl.worksheet.table import Table

from app.analysis import analyze


def show(label, wb, **kw):
    b = io.BytesIO(); wb.save(b)
    try:
        r = analyze(b.getvalue(), "x.xlsx", **kw)
        print("OK ", label, "| rows", r["rows"], "| status", r["executive"]["status"], "|", [i["title"] for i in r["executive"]["issues"]])
        return r
    except Exception as e:
        print("ERR", label, str(e)[:120])

# merged cells
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["region", "team", "amount"])
ws.append(["North", "a", 10]); ws.append([None, "b", 20]); ws.append([None, "c", 30]); ws.append(["South", "a", 5])
ws.merge_cells("A2:A4")
show("merged region cells", wb)

# hidden rows / columns
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["id", "secret", "amount"])
for i in range(5): ws.append([i, f"s{i}", i * 10])
ws.column_dimensions["B"].hidden = True
ws.row_dimensions[3].hidden = True
r = show("hidden col+row", wb)

# hidden sheet + autofilter + table
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Visible"
ws.append(["id", "v"]); [ws.append([i, i]) for i in range(4)]
ws.auto_filter.ref = "A1:B5"
hidden = wb.create_sheet("Hidden"); hidden.sheet_state = "hidden"; hidden.append(["x"]); hidden.append([1])
show("hidden sheet + filter", wb)
ws2 = wb.create_sheet("Tbl"); ws2.append(["a", "b"]); ws2.append([1, 2]); ws2.append([3, 4]); ws2.add_table(Table(displayName="T1", ref="A1:B3"))
show("table sheet", wb, sheet="Tbl")

# numbers stored as text
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["id", "qty", "price"])
for i in range(1, 9): ws.append([i, str(i * 3), str(i * 10.5)])
r = show("numbers stored as text", wb)

# dates stored as text, mixed
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["id", "when"])
for i in range(1, 6): ws.append([i, f"2024-0{i}-15"])
ws.append([6, datetime.date(2024, 6, 15)])
show("dates as text + real", wb)

# title rows, then table, then totals row
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Monthly report", None, None]); ws.append(["Prepared by finance", None, None]); ws.append([None, None, None])
ws.append(["id", "team", "amount"])
for i in range(1, 7): ws.append([i, "ab"[i % 2], i * 100])
ws.append(["Total", None, 2100])
r = show("title + totals row", wb)
if r: print("   header suggestion:", r["provenance"]["header_row_suggestion"], "header_row:", r["provenance"]["header_row"])

# formulas with cached values absent (openpyxl-written) vs data_only
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["a", "b", "sum"])
for i in range(1, 5): ws.append([i, i * 2, f"=A{i+1}+B{i+1}"])
show("formulas never calculated", wb)

# leading-zero ids as text and as numbers
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["code", "n"]); ws.append(["00123", 1]); ws.append(["00456", 2]); ws.append(["00789", 3])
r = show("leading zeros text", wb)
if r: print("   preview:", r["preview"][0])

# percent / currency formatted numbers
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["id", "rate", "price"])
for i in range(1, 5):
    ws.append([i, 0.1 * i, 1000 * i]); ws.cell(row=i + 1, column=2).number_format = "0%"; ws.cell(row=i + 1, column=3).number_format = '"฿"#,##0.00'
r = show("formatted numbers", wb)
if r: print("   rate stats:", r["columns"][1]["stats"]["max"], "price:", r["columns"][2]["stats"]["max"])
