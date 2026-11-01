import datetime
import io

import openpyxl

from app.analysis import analyze


def try_(label, fn):
    try:
        r = fn()
        print("OK  ", label, "->", r)
    except Exception as exc:
        print("ERR ", label, "->", str(exc)[:110])

try_("trailing comma in every row", lambda: (lambda r: (r["column_count"], [c["name"] for c in r["columns"]], r["rows"]))(analyze(b"id,v,\n1,2,\n3,4,\n5,6,\n", "t.csv")))
try_("duplicate header names", lambda: (lambda r: [c["name"] for c in r["columns"]])(analyze(b"id,total,total\n1,2,3\n4,5,6\n", "d.csv")))
try_("blank header cell mid", lambda: (lambda r: [c["name"] for c in r["columns"]])(analyze(b"id,,v\n1,x,3\n4,y,6\n", "b.csv")))
try_("spaces in header", lambda: (lambda r: [c["name"] for c in r["columns"]])(analyze(b" id , amount \n1,2\n3,4\n", "s.csv")))
try_("case-different headers", lambda: (lambda r: [c["name"] for c in r["columns"]])(analyze(b"Name,name\na,b\nc,d\n", "c.csv")))
try_("numbers as text 60/40", lambda: (lambda r: (r["columns"][1]["type"], r["columns"][1]["mismatches"]))(analyze(("id,v\n" + "\n".join(f"{i},{'abc' if i % 5 < 2 else i}" for i in range(50)) + "\n").encode(), "x.csv")))
try_("leading zeros kept", lambda: analyze(b"code,amt\n00123,1\n00456,2\n00789,3\n", "z.csv")["preview"][0])
try_("all day<=12 slash dates", lambda: (lambda r: (r["provenance"]["date_order_ambiguous"], r["provenance"]["date_order"]))(analyze(b"id,d\n1,01/02/2024\n2,03/04/2024\n3,05/06/2024\n", "a.csv")))
try_("header text-only table", lambda: (lambda r: (r["header"]["used"], [c["name"] for c in r["columns"]]))(analyze(b"name,city\nAnna,Rome\nBen,Oslo\nCy,Paris\n", "h.csv")))
try_("no header text-only table", lambda: (lambda r: (r["header"]["used"], r["rows"]))(analyze(b"Anna,Rome\nBen,Oslo\nCy,Paris\n", "h2.csv")))
try_("thai header", lambda: (lambda r: [c["name"] for c in r["columns"]])(analyze("ชื่อ,เมือง,ยอด\nสมชาย,กรุงเทพ,100\nสมหญิง,ภูเก็ต,200\n".encode(), "t.csv")))
try_("semicolon + decimal comma", lambda: (lambda r: (r["column_count"], r["columns"][1]["type"], r["columns"][1]["stats"]["mean"]))(analyze(b"id;amt\n1;1,5\n2;2,5\n3;3,5\n", "e.csv")))
try_("1,234 ambiguous", lambda: (lambda r: (r["columns"][1]["type"], r["columns"][1]["stats"]["mean"]))(analyze(b'id,amt\n1,"1,234"\n2,"2,500"\n3,"3,100"\n', "k.csv")))
try_("percent column", lambda: (lambda r: (r["columns"][1]["type"], r["columns"][1].get("format"), r["columns"][1]["stats"]["max"]))(analyze(b"id,rate\n1,15%\n2,20%\n3,7.5%\n", "p.csv")))
try_("boolean variants", lambda: (lambda r: (r["columns"][1]["type"], r["columns"][1]["mismatches"]))(analyze(b"id,ok\n1,true\n2,FALSE\n3,True\n4,yes\n", "bo.csv")))
try_("yes/no column", lambda: (lambda r: (r["columns"][1]["type"]))(analyze(b"id,ok\n1,yes\n2,no\n3,yes\n4,no\n", "yn.csv")))
try_("very long text cell", lambda: analyze(("id,t\n1," + "x" * 100000 + "\n2,y\n").encode(), "l.csv")["rows"])
try_("single column with blanks", lambda: analyze(b"v\n1\n\n3\n", "sc.csv"))
try_("xml-looking text", lambda: analyze(b"id,t\n1,<b>hi</b>\n2,a&b\n", "x.csv")["rows"])
try_("formula-like text", lambda: analyze(b"id,t\n1,=SUM(A1)\n2,+cmd\n3,@x\n", "f.csv")["rows"])
try_("only header", lambda: analyze(b"a,b\n", "oh.csv"))
try_("empty file", lambda: analyze(b"", "e.csv"))
try_("binary garbage", lambda: analyze(bytes(range(256)) * 10, "g.csv"))
try_("huge numbers", lambda: (lambda r: (r["columns"][1]["type"], r["columns"][1]["stats"]["max"]))(analyze(b"id,v\n1,1e300\n2,5\n3,7\n", "hn.csv")))
try_("inf nan words", lambda: (lambda r: (r["columns"][1]["type"], r["columns"][1]["invalid"]))(analyze(b"id,v\n1,inf\n2,nan\n3,5\n", "in.csv")))

# xlsx edge cases
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["id", "name", "amt", "when"]); ws.append([1, "a", "100", datetime.date(2024, 1, 2)])
ws.append([None, None, None, None])
ws.append([2, "b", 200, datetime.date(2024, 1, 3)]); ws.append([3, None, "x", None])
ws.merge_cells("A6:B6"); ws["A6"] = "merged"
b = io.BytesIO(); wb.save(b)
try_("xlsx blank row + text numbers", lambda: (lambda r: (r["rows"], r["columns"][2]["type"], r["columns"][2]["mismatches"], r["issue_samples"]["kinds"].get("missing", {}).get("examples")))(analyze(b.getvalue(), "x.xlsx")))
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["id", None, "v"]); ws.append([1, "x", 2]); ws.append([2, "y", 3])
b = io.BytesIO(); wb.save(b)
try_("xlsx blank header cell", lambda: [c["name"] for c in analyze(b.getvalue(), "x.xlsx")["columns"]])
