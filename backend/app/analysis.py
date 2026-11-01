"""Deterministic CSV profiling. All scores describe inferred rules, not truth."""

import codecs
import csv
import datetime
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from functools import lru_cache

import numpy as np
import openpyxl
import pandas as pd

MAX_BYTES = 40 * 1024 * 1024
MAX_CELLS = 4_000_000
MAX_ROWS = 100_000
MAX_COLUMNS = 100
MAX_XLSX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_MISSING_MARKERS = 20
MAX_SAMPLES = 1000
JSON_SUFFIXES = (".json", ".jsonl", ".ndjson")
ALLOWED_TYPES = {"numeric", "categorical", "datetime", "boolean", "email", "identifier"}
EMAIL_PATTERN = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


DATE_YMD = re.compile(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[T ]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})?)?")
DATE_DMY = re.compile(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})(?:[T ]\d{1,2}:\d{2}(?::\d{2})?)?")


CURRENCY_PREFIX = re.compile(r"(?:฿|\$|€|£|¥|THB)\s*", re.IGNORECASE)
CURRENCY_SUFFIX = re.compile(r"\s*(?:บาท|฿|THB)$", re.IGNORECASE)
PLAIN_NUMBER = re.compile(
    r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d*)?(?:[eE][+-]?\d+)?|[+-]?\.\d+(?:[eE][+-]?\d+)?|[+-]?(?:inf(?:inity)?|nan)",
    re.IGNORECASE,
)


EURO_NUMBER = re.compile(r"[+-]?(?:\d{1,3}(?:[. \u00a0]\d{3})+|\d+),(\d+)")


def _euro_number(text: str) -> float | None:
    """1.234,50 / 1 234,50 / 12,5 -> float. A lone comma followed by exactly three digits is English thousands."""
    match = EURO_NUMBER.fullmatch(text)
    if not match:
        return None
    integer_part = text[: match.start(1) - 1]
    grouped = bool(re.search(r"[. \u00a0]\d{3}", integer_part))
    if not grouped and len(match.group(1)) == 3:
        return None
    return float(re.sub(r"[. \u00a0]", "", integer_part) + "." + match.group(1))


@lru_cache(maxsize=200_000)
def _parse_number(value: str) -> tuple[float | None, str | None]:
    """Parse plain, thousands-separated, currency, percent and accounting-negative numbers.

    Returns (number, format). Percent values keep their face value (15% -> 15.0)."""
    text = value.strip()
    if PLAIN_NUMBER.fullmatch(text):
        return float(text.replace(",", "")), None
    if (european := _euro_number(text)) is not None:
        return european, None
    negative = False
    if len(text) > 2 and text[0] == "(" and text[-1] == ")":
        negative, text = True, text[1:-1].strip()
    sign = ""
    if text[:1] in ("+", "-"):
        sign, text = text[0], text[1:].lstrip()
    fmt = None
    if match := CURRENCY_PREFIX.match(text):
        text, fmt = text[match.end() :], "currency"
    if match := CURRENCY_SUFFIX.search(text):
        text, fmt = text[: match.start()], "currency"
    if text.endswith("%"):
        text, fmt = text[:-1].rstrip(), "percent"
    if sign and text[:1] in ("+", "-"):
        return None, None
    if not PLAIN_NUMBER.fullmatch(text):
        if (european := _euro_number(text)) is not None:
            return (-european if negative != (sign == "-") else european), fmt
        return None, None
    number = float((sign + text).replace(",", ""))
    if fmt and not np.isfinite(number):
        return None, None
    return (-number if negative else number), fmt


def _map_unique(series: pd.Series, function) -> pd.Series:
    """Apply a per-value function once per distinct value; repetitive columns are far cheaper this way."""
    mapping = {value: function(value) for value in series.dropna().unique()}
    return series.map(mapping)


def _numeric_series(series: pd.Series) -> pd.Series:
    return _map_unique(series, lambda cell: _parse_number(cell)[0]).astype(float)


MONTH_NAMES = {
    **{name: number for number, name in enumerate(
        ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"], 1)},
    **{name: number for number, name in enumerate(
        ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)},
    "sept": 9,
    **{name: number for number, name in enumerate(
        ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"], 1)},
    **{name: number for number, name in enumerate(
        ["ม.ค", "ก.พ", "มี.ค", "เม.ย", "พ.ค", "มิ.ย", "ก.ค", "ส.ค", "ก.ย", "ต.ค", "พ.ย", "ธ.ค"], 1)},
}
MONTH_TOKEN = r"[^\s\d/\-.,]+(?:\.[^\s\d/\-.,]+)*\.?"
DATE_DAY_MONTH_NAME = re.compile(rf"(\d{{1,2}})[ \-/.]*({MONTH_TOKEN})[ \-/.,]*(\d{{4}})")
DATE_MONTH_NAME_DAY = re.compile(rf"({MONTH_TOKEN})[ .]+(\d{{1,2}}),?[ ]+(\d{{4}})")


TWO_DIGIT_YEAR_DATE = re.compile(r"(\d{1,2})([-/.])(\d{1,2})\2(\d{2})((?:[T ]\d{1,2}:\d{2}(?::\d{2})?)?)")


def _expand_year(value: str, era: str) -> str:
    """D/M/YY -> D/M/YYYY. Buddhist Era reads YY as 25YY; Christian Era reads 00-69 as 20YY and 70-99 as 19YY."""
    match = TWO_DIGIT_YEAR_DATE.fullmatch(value)
    if not match:
        return value
    short = int(match[4])
    year = 2500 + short if era == "be" else (2000 + short if short <= 69 else 1900 + short)
    return f"{match[1]}{match[2]}{match[3]}{match[2]}{year}{match[5]}"


def _month_number(token: str) -> int | None:
    return MONTH_NAMES.get(token.strip().lower().rstrip("."))


def _date_parts(value: str, order: str = "dmy") -> tuple[int, int, int] | None:
    """Return (year, month, day) for a date-shaped value; D/M/YYYY unless order is "mdy". None when not date-shaped."""
    if match := DATE_YMD.fullmatch(value):
        return int(match[1]), int(match[2]), int(match[3])
    if (match := DATE_DAY_MONTH_NAME.fullmatch(value)) and (month := _month_number(match[2])) is not None:
        return int(match[3]), month, int(match[1])
    if (match := DATE_MONTH_NAME_DAY.fullmatch(value)) and (month := _month_number(match[1])) is not None:
        return int(match[3]), month, int(match[2])
    if match := DATE_DMY.fullmatch(value):
        first, second = int(match[1]), int(match[2])
        return (int(match[3]), second, first) if order == "dmy" else (int(match[3]), first, second)
    return None


def _to_date(value: str, order: str = "dmy") -> datetime.date | None:
    """Calendar date for a date-shaped value; Buddhist Era years (2400+) are converted by subtracting 543."""
    parts = _date_parts(value, order)
    if parts is None:
        return None
    year, month, day = parts
    if year >= 2400:
        year -= 543
    try:
        return datetime.date(year, month, day)
    except ValueError:
        return None


def _is_valid_date(value: str, order: str = "dmy") -> bool:
    return _to_date(value, order) is not None


AMOUNT_NAME = re.compile(r"amount|price|total|qty|quantity|revenue|cost|salary|ราคา|จำนวน|ยอด|เงิน", re.IGNORECASE)
PAST_DATE_NAME = re.compile(r"order|created|signup|sign_up|birth|dob|joined|sent|paid|invoice|delivered|สั่ง|เกิด|สร้าง|ออก|ส่ง", re.IGNORECASE)
SERIAL_DATE_NAME = re.compile(r"date|time|due|deadline|created|updated|birth|dob|วันที่|วัน|กำหนด", re.IGNORECASE)
THAI_ID_PATTERN = re.compile(r"\d{13}")
THAI_PHONE_PATTERN = re.compile(r"0[2-9]\d{7,8}")
THAI_ID_NAME_HINTS = ("national_id", "citizen", "id_card", "idcard", "thai_id", "เลขบัตร", "บัตรประชาชน")


def _thai_id_valid(value: str) -> bool:
    """Thai national ID check digit: (11 - sum(d_i * (13 - i)) mod 11) mod 10."""
    if not THAI_ID_PATTERN.fullmatch(value) or value[0] == "0":
        return False
    total = sum(int(digit) * (13 - index) for index, digit in enumerate(value[:12]))
    return (11 - total % 11) % 10 == int(value[12])


def _detect_personal_pattern(name: str, values: pd.Series) -> str | None:
    """Flag columns that look like Thai national IDs or phone numbers (personal data)."""
    if not len(values):
        return None
    lowered = name.lower()
    counts = values.value_counts()
    weights = counts.to_numpy()
    total = weights.sum()
    compact = [re.sub(r"[\s-]", "", value) for value in counts.index]
    id_shaped = np.array([bool(THAI_ID_PATTERN.fullmatch(value)) for value in compact])
    if weights[id_shaped].sum() / total >= 0.8:
        valid = np.array([_thai_id_valid(value) for value in compact])
        valid_share = weights[id_shaped & valid].sum() / weights[id_shaped].sum()
        if valid_share >= 0.5 or any(hint in lowered for hint in THAI_ID_NAME_HINTS):
            return "thai_national_id"
    phone_shaped = np.array([bool(THAI_PHONE_PATTERN.fullmatch(value)) for value in compact])
    if weights[phone_shaped].sum() / total >= 0.8:
        return "thai_phone"
    return None


def _timeseries(column: str, values: pd.Series, order: str = "dmy") -> dict | None:
    """Row counts per day, month or year for one date column, with empty periods counted as gaps."""
    parsed = [parsed_date for parsed_date in (_to_date(value, order) for value in values) if parsed_date]
    if not parsed:
        return None
    first, last = min(parsed), max(parsed)
    span = (last - first).days
    if span <= 92:
        granularity, key = "day", lambda d: d.isoformat()
        labels = [(first + datetime.timedelta(days=offset)).isoformat() for offset in range(span + 1)]
    elif span <= 3650:
        granularity, key = "month", lambda d: f"{d.year:04d}-{d.month:02d}"
        labels = []
        year, month = first.year, first.month
        while (year, month) <= (last.year, last.month):
            labels.append(f"{year:04d}-{month:02d}")
            year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    else:
        granularity, key = "year", lambda d: f"{d.year:04d}"
        labels = [f"{year:04d}" for year in range(first.year, last.year + 1)]
    counts = Counter(key(d) for d in parsed)
    points = [{"label": label, "count": counts.get(label, 0)} for label in labels]
    return {
        "column": column,
        "granularity": granularity,
        "first": first.isoformat(),
        "last": last.isoformat(),
        "gaps": sum(point["count"] == 0 for point in points),
        "points": points,
    }


def _xlsx_to_csv_text(content: bytes, sheet: str | None = None) -> tuple[str, str, list[str]]:
    """Read cached values of the first worksheet. Formulas, macros and external links are never evaluated."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_XLSX_UNCOMPRESSED:
                raise ValueError("The workbook is too large after decompression.")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("The Excel workbook could not be read.") from exc
    except ValueError:
        raise
    except Exception as exc:  # openpyxl raises several unrelated types for corrupt workbooks
        raise ValueError("The Excel workbook could not be read.") from exc
    try:
        names = list(workbook.sheetnames)
        if sheet is not None and sheet not in names:
            raise ValueError("The selected worksheet does not exist in this workbook.")
        worksheet = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            cells = [_xlsx_cell(cell) for cell in row]
            while cells and cells[-1] == "":
                cells.pop()
            rows.append(cells)
            if len(rows) > MAX_ROWS + 2:
                raise ValueError("Maximum 100,000 rows supported.")
        while rows and not rows[-1]:
            rows.pop()
        width = max((len(row) for row in rows), default=0)
        if not width:
            raise ValueError(f"Worksheet “{worksheet.title}” is empty.")
        if width > MAX_COLUMNS:
            raise ValueError("Maximum 100 columns supported.")
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        for row in rows:
            writer.writerow(row + [""] * (width - len(row)))
        return buffer.getvalue(), worksheet.title, names
    finally:
        workbook.close()


def _xlsx_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _flatten(value, prefix: str = "", depth: int = 0) -> dict:
    """Flatten nested objects into dotted columns; arrays and very deep objects become JSON text."""
    if isinstance(value, dict) and depth < 3 and value:
        flat = {}
        for key, item in value.items():
            flat.update(_flatten(item, f"{prefix}.{key}" if prefix else str(key), depth + 1))
        return flat
    return {prefix or "value": value}


def _json_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _json_to_csv_text(content: bytes) -> tuple[str, str]:
    """Turn an array of records (or JSON Lines) into CSV text with the union of keys as columns."""
    raw, encoding = _decode(content)
    try:
        try:
            document = json.loads(raw)
        except json.JSONDecodeError:
            lines = [line for line in raw.splitlines() if line.strip()]
            document = [json.loads(line) for line in lines]
    except (json.JSONDecodeError, RecursionError) as exc:
        raise ValueError("The JSON file is not valid JSON or JSON Lines.") from exc
    if isinstance(document, dict):
        candidates = [
            key for key, item in document.items() if isinstance(item, list) and item and all(isinstance(r, dict) for r in item)
        ]
        if len(candidates) == 1:
            document = document[candidates[0]]
        elif not candidates:
            document = [document]
        else:
            raise ValueError(f"The JSON has several record arrays ({', '.join(candidates[:5])}); keep one.")
    if not isinstance(document, list) or not document:
        raise ValueError("The JSON must contain a non-empty array of records.")
    if len(document) > MAX_ROWS:
        raise ValueError("Maximum 100,000 rows supported.")
    records = [_flatten(item) if isinstance(item, dict) else {"value": item} for item in document]
    keys: dict[str, None] = {}
    for record in records:
        keys.update(dict.fromkeys(record))
    if len(keys) > MAX_COLUMNS:
        raise ValueError("Maximum 100 columns supported.")
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(list(keys))
    for record in records:
        writer.writerow([_json_cell(record.get(key)) for key in keys])
    return buffer.getvalue(), encoding


DELIMITERS = {",": "comma", ";": "semicolon", "	": "tab", "|": "pipe"}


def _decode(content: bytes) -> tuple[str, str]:
    """Decode deterministically: BOM, then UTF-8, then Thai Windows-874 (common for Excel exports)."""
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        try:
            return content.decode("utf-16"), "utf-16"
        except UnicodeDecodeError as exc:
            raise ValueError("The UTF-16 file could not be decoded.") from exc
    try:
        return content.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return content.decode("cp874"), "cp874"
    except UnicodeDecodeError as exc:
        raise ValueError("Use a UTF-8, UTF-16 or Thai Windows-874 encoded CSV file.") from exc


def _detect_delimiter(text: str) -> str:
    """Pick the delimiter that splits the first rows into the same number (>1) of fields; comma wins ties."""
    best, best_score = ",", 0
    for candidate in DELIMITERS:
        sample = []
        for row in csv.reader(io.StringIO(text), delimiter=candidate):
            sample.append(len(row))
            if len(sample) == 20:
                break
        if sample and sample[0] > 1 and len(set(sample)) == 1 and sample[0] > best_score:
            best, best_score = candidate, sample[0]
    return best


def _detect_header(first: list[str], second: list[str]) -> bool:
    """Treat row one as a header only when its contents provide clear evidence."""
    common_labels = {
        "id",
        "name",
        "first_name",
        "last_name",
        "email",
        "date",
        "time",
        "amount",
        "value",
        "group",
        "category",
        "country",
        "region",
        "status",
        "type",
        "description",
        "quantity",
        "price",
        "revenue",
    }
    normalized = [cell.strip().lower().replace(" ", "_") for cell in first]
    label_shaped = all(
        cell
        and (cell[0].isalpha() or cell[0] == "_")
        and all(char.isalnum() or char in "_-" or unicodedata.category(char).startswith("M") for char in cell)
        for cell in normalized
    )
    visibly_header_shaped = label_shaped and all(cell.strip() == cell.strip().lower() for cell in first)
    data_shaped_second_row = any(
        re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", cell.strip())
        or _date_parts(cell.strip()) is not None
        for cell in second
    )
    return bool(
        any(cell in common_labels for cell in normalized)
        or (visibly_header_shaped and (data_shaped_second_row or not second or all(not cell.strip() for cell in second)))
    )


MAX_PROBLEM_ROWS = 500
MAX_PROBLEM_CELLS = 60_000


def _problem_rows(frame: pd.DataFrame, positions: set[int], samples: dict, row_offset: int) -> dict:
    """The original rows behind the located issues (lowest row numbers first), so they can be handed to whoever fixes them."""
    labels: dict[int, list[str]] = {}
    for kind, entry in samples.items():
        for example in entry["examples"]:
            found = labels.setdefault(example["row"], [])
            label = f"{kind}:{example['column']}" if example["column"] else kind
            if label not in found:
                found.append(label)
    limit = max(1, min(MAX_PROBLEM_ROWS, MAX_PROBLEM_CELLS // max(len(frame.columns), 1)))
    chosen = sorted(positions)[:limit]
    return {
        "total": len(positions),
        "listed": len(chosen),
        "columns": list(frame.columns),
        "rows": [
            {
                "row": position + row_offset,
                "issues": labels.get(position + row_offset, []),
                "values": [value[:120] for value in frame.iloc[position].tolist()],
            }
            for position in chosen
        ],
    }


def _executive_summary(
    *,
    scores: dict,
    missing: int,
    duplicates: int,
    rows: int,
    columns: int,
    mismatches: int,
    invalid: int,
    header_detected: bool,
    outliers: int,
    correlations: list[dict],
    embedded_delimiters: list[str],
    business_rule_violations: int,
    personal_data: list[str],
    key_columns: list[str],
    key_duplicates: int,
    serial_date_columns: list[str],
    mismatch_examples: list[tuple[str, str, list[str]]],
    context_warnings: list[dict],
    variant_columns: list[tuple[str, int, str]],
) -> dict:
    """Build the executive summary in English and Thai; the English `status` stays canonical for workflow checks."""
    issues = []

    def add(severity, title, detail, recommendation):
        issues.append({"severity": severity, "title": title, "detail": detail, "recommendation": recommendation})

    if not header_detected:
        add(
            "medium",
            ("Column names are missing", "ไม่พบชื่อคอลัมน์"),
            (
                "The first row appears to contain data, so neutral column names were generated.",
                "แถวแรกดูเหมือนเป็นข้อมูล จึงสร้างชื่อคอลัมน์กลางให้แทน",
            ),
            (
                "Rename columns before sharing or connecting this dataset to reporting tools.",
                "ตั้งชื่อคอลัมน์ก่อนส่งต่อหรือเชื่อมต่อกับเครื่องมือรายงาน",
            ),
        )
    if missing:
        ratio = missing / max(1, rows * columns)
        add(
            "high" if ratio >= 0.05 else "medium",
            (f"{missing:,} missing cells need review", f"มี {missing:,} ช่องที่ว่างและต้องตรวจสอบ"),
            (
                "Blank values can change totals, averages, and segment counts.",
                "ค่าว่างอาจทำให้ผลรวม ค่าเฉลี่ย และจำนวนในแต่ละกลุ่มคลาดเคลื่อน",
            ),
            (
                "Confirm whether blanks mean unknown, not applicable, or a data collection failure.",
                "ยืนยันว่าค่าว่างหมายถึงไม่ทราบ ไม่เกี่ยวข้อง หรือเก็บข้อมูลไม่สำเร็จ",
            ),
        )
    if duplicates:
        add(
            "high" if duplicates / rows >= 0.02 else "medium",
            (f"{duplicates:,} duplicate rows may overstate results", f"มี {duplicates:,} แถวซ้ำที่อาจทำให้ผลลัพธ์สูงเกินจริง"),
            (
                "Repeated records can inflate counts and monetary totals.",
                "ระเบียนที่ซ้ำอาจทำให้จำนวนและยอดเงินสูงเกินจริง",
            ),
            (
                "Confirm the business key, then remove only records proven to be duplicates.",
                "ยืนยันคีย์ทางธุรกิจ แล้วลบเฉพาะระเบียนที่พิสูจน์แล้วว่าซ้ำ",
            ),
        )
    if mismatches:
        add(
            "medium",
            (f"{mismatches:,} values conflict with their column type", f"มี {mismatches:,} ค่าที่ขัดกับชนิดของคอลัมน์"),
            (
                "Mixed types can break sorting, aggregation, and downstream imports."
                + "".join(f" In {name} (expected {dtype}): {', '.join(repr(v) for v in found)}." for name, dtype, found in mismatch_examples[:3]),
                "ชนิดข้อมูลที่ปนกันอาจทำให้การเรียงลำดับ การรวมยอด และการนำเข้าระบบอื่นผิดพลาด"
                + "".join(f" ใน {name} (คาดว่าเป็น {dtype}): {', '.join(repr(v) for v in found)}" for name, dtype, found in mismatch_examples[:3]),
            ),
            ("Standardize the flagged values before analysis.", "ปรับค่าที่ถูกทำเครื่องหมายให้เป็นรูปแบบเดียวกันก่อนวิเคราะห์"),
        )
    if invalid:
        add(
            "high",
            (f"{invalid:,} typed values are invalid", f"มี {invalid:,} ค่าที่ไม่ถูกต้องตามชนิดข้อมูล"),
            (
                "These values match the expected shape but fail a basic validity check.",
                "ค่าเหล่านี้มีรูปแบบตามที่คาดไว้ แต่ไม่ผ่านการตรวจความถูกต้องพื้นฐาน",
            ),
            (
                "Correct or exclude the invalid values with an auditable rule.",
                "แก้ไขหรือตัดค่าที่ไม่ถูกต้องออกด้วยกฎที่ตรวจสอบย้อนหลังได้",
            ),
        )
    if outliers:
        add(
            "info",
            (f"{outliers:,} statistical outliers deserve context", f"มี {outliers:,} ค่าผิดปกติทางสถิติที่ควรหาสาเหตุ"),
            (
                "IQR flags unusual numeric values; unusual does not automatically mean incorrect.",
                "IQR ระบุค่าตัวเลขที่ผิดปกติ แต่ค่าผิดปกติไม่ได้แปลว่าผิดเสมอไป",
            ),
            (
                "Ask the data owner whether these represent valid edge cases, errors, or exceptional events.",
                "สอบถามเจ้าของข้อมูลว่าเป็นกรณีสุดโต่งที่ถูกต้อง ข้อผิดพลาด หรือเหตุการณ์พิเศษ",
            ),
        )
    if embedded_delimiters:
        names = ", ".join(embedded_delimiters[:3])
        add(
            "medium",
            ("Possible fields embedded inside one CSV column", "อาจมีหลายฟิลด์ถูกรวมอยู่ในคอลัมน์เดียว"),
            (f"Repeated pipe delimiters were found in: {names}.", f"พบตัวคั่น | ซ้ำหลายตัวใน: {names}"),
            (
                "Confirm whether each pipe-separated field needs its own CSV column before reporting.",
                "ยืนยันว่าแต่ละฟิลด์ที่คั่นด้วย | ควรแยกเป็นคอลัมน์ของตัวเองก่อนจัดทำรายงาน",
            ),
        )
    if business_rule_violations:
        add(
            "high",
            (
                f"{business_rule_violations:,} business-rule violations require action",
                f"มี {business_rule_violations:,} รายการที่ละเมิดกฎทางธุรกิจและต้องดำเนินการ",
            ),
            (
                "Uploaded values fall outside rules explicitly configured for this analysis.",
                "ค่าในไฟล์อยู่นอกกฎที่กำหนดไว้สำหรับการวิเคราะห์นี้",
            ),
            (
                "Resolve each violation or document an approved exception before executive use.",
                "แก้ไขแต่ละรายการ หรือบันทึกข้อยกเว้นที่ได้รับอนุมัติ ก่อนนำเสนอผู้บริหาร",
            ),
        )
    if key_duplicates:
        add(
            "high",
            (
                f"{key_duplicates:,} rows repeat the business key ({', '.join(key_columns)})",
                f"มี {key_duplicates:,} แถวที่คีย์ทางธุรกิจซ้ำ ({', '.join(key_columns)})",
            ),
            (
                "The same key appears more than once, so counts and totals per key may be overstated.",
                "คีย์เดียวกันปรากฏมากกว่าหนึ่งครั้ง จำนวนและยอดรวมต่อคีย์อาจสูงเกินจริง",
            ),
            (
                "Confirm which record is authoritative for each key and remove or merge the others.",
                "ยืนยันว่าระเบียนใดเป็นข้อมูลหลักของแต่ละคีย์ แล้วลบหรือรวมระเบียนที่เหลือ",
            ),
        )
    if personal_data:
        add(
            "high",
            ("Possible personal data detected", "ตรวจพบข้อมูลที่อาจเป็นข้อมูลส่วนบุคคล"),
            (
                f"Thai national ID or phone number patterns were found in: {', '.join(personal_data[:3])}.",
                f"พบรูปแบบเลขบัตรประชาชนหรือเบอร์โทรศัพท์ไทยใน: {', '.join(personal_data[:3])}",
            ),
            (
                "Remove or mask these columns, or confirm a lawful basis and approved handling before sharing.",
                "ลบหรือปิดบังคอลัมน์เหล่านี้ หรือยืนยันฐานทางกฎหมายและวิธีจัดการที่ได้รับอนุมัติก่อนส่งต่อ",
            ),
        )
    for warning in context_warnings:
        if warning["kind"] == "future_date":
            add(
                "medium",
                (f"{warning['count']:,} dates in {warning['column']} are in the future", f"มี {warning['count']:,} วันที่ใน {warning['column']} ที่อยู่ในอนาคต"),
                (
                    "A column named like a past event holds dates later than today.",
                    "คอลัมน์ที่ชื่อบ่งว่าเป็นเหตุการณ์ในอดีตมีวันที่หลังวันนี้",
                ),
                ("Check the year or the day/month order of these dates.", "ตรวจปีหรือลำดับวัน/เดือนของวันที่เหล่านี้"),
            )
        else:
            add(
                "medium",
                (f"{warning['count']:,} negative values in {warning['column']}", f"มี {warning['count']:,} ค่าติดลบใน {warning['column']}"),
                (
                    "Amounts, prices and quantities are normally not negative; refunds or entry errors would show up this way.",
                    "ยอดเงิน ราคา และจำนวนมักไม่ติดลบ รายการคืนเงินหรือการคีย์ผิดจะปรากฏแบบนี้",
                ),
                ("Confirm each one is intended (for example a refund) or correct it.", "ยืนยันว่าตั้งใจ (เช่น การคืนเงิน) หรือแก้ไข"),
            )
    if variant_columns:
        names = ", ".join(name for name, _, _ in variant_columns[:3])
        example = variant_columns[0][2]
        add(
            "medium",
            (f"Same value spelled differently in {names}", f"ค่าเดียวกันสะกดต่างกันใน {names}"),
            (
                f"Values that differ only by capital letters or spacing are counted as different groups (for example {example}).",
                f"ค่าที่ต่างกันเพียงตัวพิมพ์ใหญ่เล็กหรือช่องว่างจะถูกนับเป็นคนละกลุ่ม (เช่น {example})",
            ),
            ("Use one spelling for each value before grouping or sending.", "ใช้การสะกดเดียวกันสำหรับแต่ละค่าก่อนจัดกลุ่มหรือส่งต่อ"),
        )
    if serial_date_columns:
        add(
            "medium",
            ("Numbers that look like Excel dates", "ตัวเลขที่น่าจะเป็นวันที่ของ Excel"),
            (
                f"{', '.join(serial_date_columns[:3])} hold whole numbers in the range of Excel date serials (for example 45366 is 15 Mar 2024), so they are read as plain numbers, not dates.",
                f"{', '.join(serial_date_columns[:3])} เก็บเป็นจำนวนเต็มในช่วงเดียวกับเลขวันที่ของ Excel (เช่น 45366 คือ 15 มี.ค. 2024) จึงถูกอ่านเป็นตัวเลขธรรมดา ไม่ใช่วันที่",
            ),
            (
                "In Excel, format the column as a date (or write it as YYYY-MM-DD) and upload again.",
                "ใน Excel ให้จัดรูปแบบคอลัมน์เป็นวันที่ (หรือเขียนเป็น YYYY-MM-DD) แล้วอัปโหลดใหม่",
            ),
        )
    high = sum(item["severity"] == "high" for item in issues)
    if high:
        status = "Action required"
        message = ("Fix the high-impact problems before you send this data.", "แก้ปัญหาที่มีผลสูงก่อนส่งข้อมูลนี้ให้ผู้อื่น")
    elif issues:
        status = "Review needed"
        message = ("Check the listed items before you send this data.", "ตรวจรายการที่แจ้งก่อนส่งข้อมูลนี้ให้ผู้อื่น")
    else:
        status = "Ready for exploration"
        message = ("No problems were found by the supported checks, so this data is ready to send.", "ไม่พบปัญหาจากการตรวจที่รองรับ ข้อมูลนี้พร้อมส่ง")
    severity_order = {"high": 0, "medium": 1, "info": 2}
    issues.sort(key=lambda issue: severity_order[issue["severity"]])
    scope = (
        "Automated structural checks only; business accuracy still requires an accountable owner.",
        "เป็นเพียงการตรวจเชิงโครงสร้างอัตโนมัติ ความถูกต้องทางธุรกิจยังต้องมีผู้รับผิดชอบยืนยัน",
    )

    def build(index: int) -> dict:
        return {
            "status": status,
            "message": message[index],
            "quality_score": round(scores["Overall"], 2),
            "issue_count": len(issues),
            "high_priority_count": high,
            "issues": [
                {
                    "severity": issue["severity"],
                    "title": issue["title"][index],
                    "detail": issue["detail"][index],
                    "recommendation": issue["recommendation"][index],
                }
                for issue in issues
            ],
            "signals": [
                (
                    f"Observed Pearson coefficient in this file: {pair['left']} ↔ {pair['right']} ({pair['coefficient']:+.2f})."
                    if index == 0
                    else f"ค่าสัมประสิทธิ์สหสัมพันธ์เพียร์สันที่พบในไฟล์นี้: {pair['left']} ↔ {pair['right']} ({pair['coefficient']:+.2f})"
                )
                for pair in correlations[:1]
            ],
            "scope_note": scope[index],
        }

    english, thai = build(0), build(1)
    english["thai"] = thai
    return english


def _read_text(content: bytes, filename: str, sheet: str | None) -> tuple[str, str | None, list[str] | None, str, str, str]:
    """Turn an uploaded CSV, workbook or JSON file into delimited text; returns text, sheet, sheets, encoding, delimiter, format."""
    if len(content) > MAX_BYTES:
        raise ValueError(f"The file exceeds the {MAX_BYTES // (1024 * 1024)} MB limit.")
    sheet_name, sheet_names = None, None
    if sheet is not None and not isinstance(sheet, str):
        raise ValueError("The worksheet name must be text.")
    if content.startswith(b"PK\x03\x04"):
        text, sheet_name, sheet_names = _xlsx_to_csv_text(content, sheet)
        return text, sheet_name, sheet_names, "xlsx", ",", "xlsx"
    if filename.lower().endswith(JSON_SUFFIXES):
        text, encoding = _json_to_csv_text(content)
        return text, None, None, encoding, ",", "json"
    text, encoding = _decode(content)
    if "\x00" in text:
        raise ValueError("The file contains binary data.")
    return text, None, None, encoding, _detect_delimiter(text), "csv"


def analyze(
    content: bytes,
    filename: str,
    header_mode: str = "auto",
    column_names: list[str] | None = None,
    type_overrides: dict[str, str] | None = None,
    business_rules: dict[str, dict] | None = None,
    missing_markers: list[str] | None = None,
    sheet: str | None = None,
    header_row: int | None = None,
    date_order: str | None = None,
    business_key: list[str] | None = None,
    group_by: dict | None = None,
    year_era: str | None = None,
) -> dict:
    text, sheet_name, sheet_names, encoding, delimiter, source_format = _read_text(content, filename, sheet)
    markers = missing_markers or []
    if (
        not isinstance(markers, list)
        or len(markers) > MAX_MISSING_MARKERS
        or not all(isinstance(item, str) and 0 < len(item.strip()) <= 30 for item in markers)
    ):
        raise ValueError("Missing markers must be a list of at most 20 short text values.")
    marker_set = {item.strip().lower() for item in markers}
    if header_mode not in {"auto", "present", "absent"}:
        raise ValueError("header_mode must be auto, present, or absent.")
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        all_rows = list(reader)
        total_file_rows = len(all_rows)
        header_row_suggestion = None
        probe = [sum(1 for cell in row if cell.strip()) for row in all_rows[:30]]
        if probe and max(probe) >= 2:
            first_full = next(index for index, count in enumerate(probe) if count >= 0.8 * max(probe))
            header_row_suggestion = first_full + 1 if first_full else None
        skipped = 0
        if (
            header_row is None
            and header_row_suggestion
            and len({len(row) for row in all_rows}) > 1
            and len(all_rows) - header_row_suggestion >= 2
            and len({len(row) for row in all_rows[header_row_suggestion - 1 :]}) == 1
        ):
            # Title lines above the table make the rows uneven; start at the first full row instead of failing.
            header_row = header_row_suggestion
        if header_row is not None:
            if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= min(total_file_rows, 1000):
                raise ValueError("The header row must be a row number inside the file.")
            skipped = header_row - 1
            all_rows = all_rows[skipped:]
        if not all_rows:
            raise ValueError("The CSV is empty or malformed.")
        width = len(all_rows[0])
        if width > MAX_COLUMNS:
            raise ValueError("Maximum 100 columns supported.")
        uneven = [(index + 1 + skipped, len(row)) for index, row in enumerate(all_rows) if len(row) != width]
        if uneven:
            examples = ", ".join(f"row {number} has {count}" for number, count in uneven[:5])
            more = f" (+{len(uneven) - 5} more)" if len(uneven) > 5 else ""
            raise ValueError(
                f"Every row must have the same number of fields: expected {width}, but {examples}{more}."
                " If the table starts lower in the file, set the header row."
            )
        detected_header = _detect_header(all_rows[0], all_rows[1] if len(all_rows) > 1 else [])
        use_header = detected_header if header_mode == "auto" else header_mode == "present"
        if use_header:
            headers = [h.strip() for h in all_rows[0]]
            rows = all_rows[1:]
        else:
            headers = [f"column_{index + 1}" for index in range(width)]
            rows = all_rows
        if column_names is not None:
            if not isinstance(column_names, list) or not all(isinstance(name, str) for name in column_names):
                raise ValueError("Column names must be a list of text values.")
            if len(column_names) != width:
                raise ValueError("The number of column names must match the CSV width.")
            headers = [name.strip() for name in column_names]
            if any(not name for name in headers):
                raise ValueError("Column names cannot be empty.")
        if use_header and (not headers or any(not h for h in headers)):
            raise ValueError("Every column needs a non-empty header.")
        if len(set(headers)) != len(headers):
            raise ValueError("Column names must be unique.")
        if len(rows) > MAX_ROWS:
            raise ValueError("Maximum 100,000 rows supported.")
        if len(rows) * width > MAX_CELLS:
            raise ValueError(f"This file has too many cells to analyze (limit {MAX_CELLS:,} = rows x columns). Remove columns or split the file.")
    except (StopIteration, csv.Error) as exc:
        raise ValueError("The CSV is empty or malformed.") from exc
    if not rows:
        raise ValueError("The CSV must contain at least one data row.")
    frame = pd.DataFrame(rows, columns=headers)
    normalized = frame.apply(lambda col: col.str.strip().replace("", None))
    if marker_set:
        normalized = normalized.apply(lambda col: col.where(~col.str.lower().isin(marker_set), None))
    if year_era not in {None, "be", "ce"}:
        raise ValueError("year_era must be be or ce.")
    two_digit_years = 0
    for header in headers:
        column = normalized[header]
        found = _map_unique(column.dropna(), lambda value: bool(TWO_DIGIT_YEAR_DATE.fullmatch(value))).astype(bool)
        two_digit_years += int(found.sum())
        if year_era and found.any():
            normalized[header] = column.map(lambda value: _expand_year(value, year_era) if isinstance(value, str) else value)
    n = len(frame)
    type_overrides = type_overrides or {}
    if not isinstance(type_overrides, dict) or not all(
        isinstance(name, str) and isinstance(value, str) for name, value in type_overrides.items()
    ):
        raise ValueError("Type overrides must map column names to supported types.")
    unknown_overrides = set(type_overrides) - set(headers)
    invalid_override_types = set(type_overrides.values()) - ALLOWED_TYPES
    if unknown_overrides or invalid_override_types:
        raise ValueError("Schema overrides contain an unknown column or type.")
    if date_order not in {None, "dmy", "mdy"}:
        raise ValueError("date_order must be dmy or mdy.")
    seen_slash_dates = day_first_proof = month_first_proof = 0
    for header in headers:
        pairs = normalized[header].dropna().str.extract(r"^(\d{1,2})[-/.](\d{1,2})[-/.]\d{4}(?:[T ].*)?$").dropna()
        if len(pairs):
            first_part, second_part = pairs[0].astype(int), pairs[1].astype(int)
            seen_slash_dates += len(pairs)
            day_first_proof += int((first_part > 12).sum())
            month_first_proof += int((second_part > 12).sum())
    date_order_ambiguous = date_order is None and seen_slash_dates > 0 and not day_first_proof and not month_first_proof
    chosen_date_order = date_order or ("mdy" if month_first_proof and not day_first_proof else "dmy")
    columns, numeric, personal_data, serial_date_columns = [], {}, [], []
    mismatch_examples: list[tuple[str, str, list[str]]] = []
    context_warnings: list[dict] = []
    variant_columns: list[tuple[str, int, str]] = []
    today = datetime.datetime.now(tz=datetime.UTC).date()
    row_offset = 1 if source_format == "json" else (2 if use_header else 1) + skipped
    samples: dict[str, dict] = {}
    problem_positions: set[int] = set()

    def sample(kind: str, index, column: str | None, reason: str) -> None:
        """Record up to MAX_SAMPLES located examples per issue kind while counting every occurrence."""
        entry = samples.setdefault(kind, {"total": 0, "examples": []})
        entry["total"] += len(index)
        problem_positions.update(int(position) for position in index)
        room = MAX_SAMPLES - len(entry["examples"])
        for position in list(index)[: max(room, 0)]:
            if column is None:
                value = " | ".join(str(cell) for cell in normalized.loc[position].dropna().head(3))
            else:
                cell = normalized.at[position, column]
                value = None if cell is None or pd.isna(cell) else str(cell)
            entry["examples"].append(
                {"row": int(position) + row_offset, "column": column, "value": value[:80] if value else value, "reason": reason}
            )

    mismatches = invalid = checked = 0
    for name in headers:
        values = normalized[name].dropna()

        def kind(value):
            if value.lower() in {"true", "false"}:
                return "boolean"
            if _parse_number(value)[0] is not None or PLAIN_NUMBER.fullmatch(value):
                return "numeric"
            if _date_parts(value, chosen_date_order) is not None:
                return "datetime"
            return "categorical"

        kinds = _map_unique(values, kind)
        counts = Counter(kinds)
        dominant = counts.most_common(1)[0][0] if counts else "empty"
        inferred = dominant if counts and counts[dominant] / len(values) > 0.5 else "categorical"
        normalized_name = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
        email_hits = int(_map_unique(values, lambda value: bool(EMAIL_PATTERN.fullmatch(value))).sum())
        if values.size and ("email" in normalized_name or email_hits / len(values) >= 0.8):
            inferred = "email"
        if normalized_name in {"id", "customer_id", "record_id", "user_id"} or normalized_name.endswith("_id"):
            inferred = "identifier"
        pattern = _detect_personal_pattern(name, values)
        if pattern and name not in type_overrides:
            inferred = "identifier"
        if pattern:
            personal_data.append(name)
        dtype = type_overrides.get(name, inferred)
        if dtype == "email" or dtype == "identifier":
            mismatch = 0
            matched = values
        else:
            mismatch = int((kinds != dtype).sum()) if dtype not in {"categorical", "empty"} else 0
            matched = values[kinds == dtype]
        bad = 0
        stats = None
        histogram = []
        numeric_format = None
        sample("missing", normalized.index[normalized[name].isna()], name, "Blank or marked as missing")
        if (
            dtype == "numeric"
            and SERIAL_DATE_NAME.search(name)
            and matched.size
            and _numeric_series(matched).between(36526, 73415).mean() >= 0.9
            and (_numeric_series(matched) % 1 == 0).mean() >= 0.9
        ):
            serial_date_columns.append(name)
        if mismatch:
            sample("mismatch", values.index[kinds != dtype], name, f"Does not look like {dtype}")
            mismatch_examples.append((name, dtype, [str(v)[:30] for v in values[kinds != dtype].drop_duplicates().head(3)]))
        if dtype == "categorical" and 1 < values.nunique() <= 500:
            folded = _map_unique(values, lambda value: re.sub(r"\s+", " ", value).casefold())
            spelled = values.groupby(folded).agg(lambda group: group.value_counts().index.tolist())
            groups = [(key, variants) for key, variants in spelled.items() if len(variants) > 1]
            for key, variants in groups[:20]:
                sample(
                    "variant",
                    values.index[(folded == key) & (values != variants[0])],
                    name,
                    f"Spelled differently from the more common “{variants[0]}”",
                )
            if groups:
                variant_columns.append((name, len(groups), f"{groups[0][1][0]} / {groups[0][1][1]}"))
        if dtype == "numeric":
            parsed = _numeric_series(matched)
            formats = Counter(_parse_number(value)[1] for value in matched)
            numeric_format = next(
                (name_ for name_ in ("percent", "currency") if matched.size and formats[name_] / matched.size >= 0.5),
                None,
            )
            good = parsed[np.isfinite(parsed)]
            bad = len(parsed) - len(good)
            negatives = int((good < 0).sum())
            if negatives and AMOUNT_NAME.search(name):
                context_warnings.append({"kind": "negative_amount", "column": name, "count": negatives})
            sample("invalid", parsed.index[~np.isfinite(parsed)], name, "Not a finite number")
            numeric[name] = _numeric_series(normalized[name]).replace([np.inf, -np.inf], np.nan)
            if len(good):
                stats = {
                    "count": len(good),
                    "min": float(good.min()),
                    "max": float(good.max()),
                    "mean": float(good.mean()),
                    "median": float(good.median()),
                    "std": float(good.std(ddof=0)),
                }
                heights, edges = np.histogram(good, bins=min(12, max(1, int(np.sqrt(len(good))))))
                histogram = [
                    {"label": f"{edges[i]:.3g}–{edges[i + 1]:.3g}", "count": int(count)}
                    for i, count in enumerate(heights)
                ]
        elif dtype == "datetime":
            invalid_dates = ~_map_unique(matched, lambda value: _is_valid_date(value, chosen_date_order)).astype(bool)
            bad = int(invalid_dates.sum())
            sample("invalid", matched.index[invalid_dates], name, "Not a valid calendar date")
            if PAST_DATE_NAME.search(name):
                dates = _map_unique(matched, lambda value: _to_date(value, chosen_date_order))
                future = int(sum(1 for item in dates if item is not None and item is not pd.NaT and item > today))
                if future:
                    context_warnings.append({"kind": "future_date", "column": name, "count": future})
        elif dtype == "email":
            invalid_emails = ~_map_unique(matched, lambda value: bool(EMAIL_PATTERN.fullmatch(value))).astype(bool)
            bad = int(invalid_emails.sum())
            sample("invalid", matched.index[invalid_emails], name, "Not a valid email shape")
        elif pattern == "thai_national_id":
            invalid_ids = ~_map_unique(matched, lambda value: _thai_id_valid(re.sub(r"[\s-]", "", value))).astype(bool)
            bad = int(invalid_ids.sum())
            sample("invalid", matched.index[invalid_ids], name, "Thai national ID check digit is wrong")
        if dtype in {"numeric", "datetime", "boolean", "email"} or pattern == "thai_national_id":
            checked += len(matched)
            invalid += bad
        mismatches += mismatch
        columns.append(
            {
                "name": name,
                "type": dtype,
                "missing": int(normalized[name].isna().sum()),
                "unique": int(values.nunique()),
                "mismatches": mismatch,
                "invalid": bad,
                "outliers": 0,
                "inferred_type": inferred,
                "pattern": pattern,
                "format": numeric_format,
                "type_overridden": name in type_overrides,
                "stats": stats,
                "histogram": histogram,
                "categories": [{"label": str(k), "count": int(v)} for k, v in values.value_counts().head(10).items()],
            }
        )
        if dtype == "numeric" and stats:
            q1, q3 = good.quantile([0.25, 0.75])
            iqr = q3 - q1
            outlier_mask = (good < q1 - 1.5 * iqr) | (good > q3 + 1.5 * iqr)
            outlier_count = int(outlier_mask.sum())
            sample("outlier", good.index[outlier_mask], name, f"Outside the 1.5×IQR range {q1 - 1.5 * iqr:,.2f} to {q3 + 1.5 * iqr:,.2f}")
            columns[-1]["outliers"] = outlier_count
            stats.update({"q1": float(q1), "q3": float(q3)})
    missing = int(normalized.isna().sum().sum())
    present = n * len(headers) - missing
    duplicate_mask = normalized.duplicated()
    duplicates = int(duplicate_mask.sum())
    sample("duplicate", normalized.index[duplicate_mask], None, "Repeats an earlier row")
    key_columns = business_key or []
    if (
        not isinstance(key_columns, list)
        or len(key_columns) > 5
        or not all(isinstance(item, str) and item in headers for item in key_columns)
        or len(set(key_columns)) != len(key_columns)
    ):
        raise ValueError("The business key must list up to 5 existing, distinct columns.")
    key_summary = None
    if key_columns:
        key_frame = normalized[key_columns]
        complete = key_frame.notna().all(axis=1)
        keyed = key_frame[complete]
        repeated = keyed.duplicated(keep="first")
        first_seen = {
            values: position
            for position, values in zip(keyed.index[~repeated], keyed[~repeated].itertuples(index=False, name=None))
        }
        entry = samples.setdefault("key_duplicate", {"total": 0, "examples": []})
        entry["total"] += int(repeated.sum())
        for position in list(keyed.index[repeated])[:MAX_SAMPLES]:
            key_values = tuple(keyed.loc[position])
            entry["examples"].append(
                {
                    "row": int(position) + row_offset,
                    "column": ", ".join(key_columns),
                    "value": " | ".join(key_values)[:80],
                    "reason": f"Repeats a key first seen at row {first_seen[key_values] + row_offset}",
                }
            )
        key_summary = {
            "columns": key_columns,
            "duplicate_rows": int(repeated.sum()),
            "blank_key_rows": int((~complete).sum()),
        }
    group_summary = None
    if group_by:
        if not isinstance(group_by, dict) or set(group_by) - {"by", "measure"} or group_by.get("by") not in headers:
            raise ValueError("The group summary needs an existing column to group by.")
        measure = group_by.get("measure")
        if measure is not None and measure not in numeric:
            raise ValueError("The group summary measure must be a numeric column.")
        grouped = pd.DataFrame(
            {"group": normalized[group_by["by"]].fillna("(blank)"), "value": numeric[measure] if measure else np.nan}
        ).groupby("group", sort=False)
        figures = grouped["value"].agg(["count", "sum", "mean", "min", "max"])
        sizes = grouped.size()
        order_key = figures["sum"] if measure else sizes
        ranked = order_key.sort_values(ascending=False, kind="stable").index
        clean = lambda number: None if pd.isna(number) else float(number)
        group_summary = {
            "by": group_by["by"],
            "measure": measure,
            "total_groups": len(sizes),
            "truncated": len(sizes) > 50,
            "rows": [
                {
                    "group": str(label),
                    "rows": int(sizes[label]),
                    "sum": clean(figures.at[label, "sum"]) if measure else None,
                    "mean": clean(figures.at[label, "mean"]) if measure else None,
                    "min": clean(figures.at[label, "min"]) if measure else None,
                    "max": clean(figures.at[label, "max"]) if measure else None,
                }
                for label in ranked[:50]
            ],
        }
    scores = {
        "Completeness": 100 * (1 - missing / (n * len(headers))),
        "Consistency": 100 * (1 - mismatches / present) if present else None,
        "Uniqueness": 100 * (1 - duplicates / n),
        "Validity": 100 * (1 - invalid / checked) if checked else None,
    }
    available = [v for v in scores.values() if v is not None]
    scores["Overall"] = sum(available) / len(available)
    numeric_frame = pd.DataFrame(numeric)
    correlations = []
    if len(numeric_frame.columns) >= 2:
        matrix = numeric_frame.corr(min_periods=3)
        for left_index, left in enumerate(matrix.columns):
            for right in matrix.columns[left_index + 1 :]:
                coefficient = matrix.loc[left, right]
                if pd.notna(coefficient):
                    correlations.append({"left": left, "right": right, "coefficient": round(float(coefficient), 4)})
        correlations.sort(key=lambda item: abs(item["coefficient"]), reverse=True)
    outliers = sum(column["outliers"] for column in columns)
    embedded_delimiters = []
    for header in headers:
        values = normalized[header].dropna().astype(str)
        pipe_ratio = float(values.str.count(r"\|").ge(2).mean()) if len(values) else 0
        if header.count("|") >= 2 or pipe_ratio >= 0.8:
            embedded_delimiters.append(header)
    if business_rules is None:
        business_rules = {}
    if not isinstance(business_rules, dict):
        raise TypeError("Business rules must be an object keyed by column name.")
    rule_results = []
    for name, rule in business_rules.items():
        if name not in headers or not isinstance(rule, dict):
            raise ValueError("Each business rule must target an existing column.")
        if set(rule) - {"required", "allowed_values", "min", "max"}:
            raise ValueError("A business rule contains an unsupported condition.")
        if "required" in rule and not isinstance(rule["required"], bool):
            raise ValueError("required business rules must be boolean values.")
        values = normalized[name]
        nonempty = values.dropna().astype(str)
        if rule.get("required") is True:
            rule_results.append({"column": name, "rule": "required", "checked": n, "violations": int(values.isna().sum())})
            sample("rule", values.index[values.isna()], name, "Required value is blank")
        if "allowed_values" in rule:
            allowed = rule["allowed_values"]
            if not isinstance(allowed, list) or len(allowed) > 100 or not all(isinstance(item, str) for item in allowed):
                raise ValueError("allowed_values must be a list of at most 100 text values.")
            allowed_set = set(allowed)
            sample("rule", nonempty.index[~nonempty.isin(allowed_set)], name, "Value is not in the allowed list")
            rule_results.append(
                {
                    "column": name,
                    "rule": "allowed_values",
                    "checked": len(nonempty),
                    "violations": int((~nonempty.isin(allowed_set)).sum()),
                }
            )
        numeric_values = _numeric_series(nonempty).replace([np.inf, -np.inf], np.nan).dropna()
        if "min" in rule and "max" in rule:
            minimum, maximum = rule["min"], rule["max"]
            if (
                isinstance(minimum, bool)
                or isinstance(maximum, bool)
                or not isinstance(minimum, (int, float))
                or not isinstance(maximum, (int, float))
                or not np.isfinite(minimum)
                or not np.isfinite(maximum)
                or minimum > maximum
            ):
                raise ValueError("Business rule minimum must be a finite number no greater than maximum.")
        for key, operator in (("min", "below_min"), ("max", "above_max")):
            if key not in rule:
                continue
            column_type = next(column["type"] for column in columns if column["name"] == name)
            if column_type != "numeric":
                raise ValueError(f"{key} business rules require a numeric column.")
            threshold = rule[key]
            if isinstance(threshold, bool) or not isinstance(threshold, (int, float)) or not np.isfinite(threshold):
                raise ValueError(f"{key} business rules must be finite numbers.")
            violation_mask = numeric_values < threshold if key == "min" else numeric_values > threshold
            violations = violation_mask.sum()
            sample(
                "rule",
                numeric_values.index[violation_mask],
                name,
                f"Below the minimum {threshold:g}" if key == "min" else f"Above the maximum {threshold:g}",
            )
            rule_results.append(
                {"column": name, "rule": operator, "threshold": float(threshold), "checked": len(numeric_values), "violations": int(violations)}
            )
    business_rule_violations = sum(item["violations"] for item in rule_results)
    executive = _executive_summary(
        scores=scores,
        missing=missing,
        duplicates=duplicates,
        rows=n,
        columns=len(headers),
        mismatches=mismatches,
        invalid=invalid,
        header_detected=use_header or column_names is not None,
        outliers=outliers,
        correlations=correlations,
        embedded_delimiters=embedded_delimiters,
        business_rule_violations=business_rule_violations,
        personal_data=personal_data,
        key_columns=key_columns,
        key_duplicates=key_summary["duplicate_rows"] if key_summary else 0,
        serial_date_columns=serial_date_columns,
        mismatch_examples=mismatch_examples,
        context_warnings=context_warnings,
        variant_columns=variant_columns,
    )
    numeric_names = list(numeric)
    scatter = []
    if len(numeric_names) >= 2:
        scatter_axes = [correlations[0]["left"], correlations[0]["right"]] if correlations else numeric_names[:2]
        pairs = pd.DataFrame({"x": numeric[scatter_axes[0]], "y": numeric[scatter_axes[1]]}).dropna()
        scatter = pairs.head(500).to_dict("records")
    else:
        scatter_axes = numeric_names[:2]
    return {
        "filename": filename,
        "rows": n,
        "column_count": len(headers),
        "missing": missing,
        "duplicates": duplicates,
        "scores": {k: round(v, 2) if v is not None else None for k, v in scores.items()},
        "header": {
            "mode": header_mode,
            "detected": detected_header,
            "used": use_header,
            "generated_names": not use_header and column_names is None,
            "configured_names": column_names is not None,
        },
        "executive": executive,
        "columns": columns,
        "preview": normalized.head(100).where(normalized.head(100).notna(), None).to_dict("records"),
        "key_duplicates": key_summary,
        "group_summary": group_summary,
        "problem_rows": _problem_rows(frame, problem_positions, samples, row_offset),
        "issue_samples": {
            "row_label": "record" if source_format == "json" else "row",
            "max_examples": MAX_SAMPLES,
            "kinds": samples,
        },
        "timeseries": next(
            (
                series
                for column in columns
                if column["type"] == "datetime"
                and (series := _timeseries(column["name"], normalized[column["name"]].dropna(), chosen_date_order))
            ),
            None,
        ),
        "scatter": scatter,
        "scatter_axes": scatter_axes,
        "correlations": correlations[:10],
        "business_rules": {
            "configured": bool(rule_results),
            "passed": business_rule_violations == 0,
            "total_violations": business_rule_violations,
            "results": rule_results,
        },
        "provenance": {
            "source": "uploaded_file",
            "sha256": hashlib.sha256(content).hexdigest(),
            "input_bytes": len(content),
            "file_rows": total_file_rows,
            "header_row": skipped + 1 if use_header else None,
            "header_row_suggestion": header_row_suggestion,
            "date_order": chosen_date_order,
            "date_order_ambiguous": date_order_ambiguous,
            "two_digit_years": two_digit_years,
            "year_era": year_era,
            "analyzed_rows": n,
            "preview_rows": min(n, 100),
            "method_version": "0.7.1",
            "source_format": source_format,
            "sheet": sheet_name,
            "sheets": sheet_names,
            "encoding": encoding,
            "delimiter": DELIMITERS[delimiter] if source_format == "csv" else None,
            "missing_markers": sorted(marker_set),
            "calculation_mode": "deterministic",
            "data_values_generated": False,
        },
    }


def redact_columns(
    content: bytes,
    filename: str,
    drop_columns: list[str],
    header_mode: str = "auto",
    sheet: str | None = None,
    header_row: int | None = None,
) -> str:
    """Return the file as CSV text without the named columns. Header handling follows the analysis settings."""
    if not isinstance(drop_columns, list) or not drop_columns or not all(isinstance(name, str) for name in drop_columns):
        raise ValueError("Choose the columns to remove.")
    text, _, _, _, delimiter, _ = _read_text(content, filename, sheet)
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
    except csv.Error as exc:
        raise ValueError("The CSV is empty or malformed.") from exc
    if header_row is not None:
        if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= min(len(rows), 1000):
            raise ValueError("The header row must be a row number inside the file.")
        rows = rows[header_row - 1 :]
    if not rows:
        raise ValueError("The CSV is empty or malformed.")
    use_header = _detect_header(rows[0], rows[1] if len(rows) > 1 else []) if header_mode == "auto" else header_mode == "present"
    names = [cell.strip() for cell in rows[0]] if use_header else [f"column_{index + 1}" for index in range(len(rows[0]))]
    unknown = set(drop_columns) - set(names)
    if unknown:
        raise ValueError("A column to remove is not in this file.")
    keep = [index for index, name in enumerate(names) if name not in set(drop_columns)]
    if not keep:
        raise ValueError("Every column would be removed.")
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    for row in rows:
        writer.writerow([row[index] if index < len(row) else "" for index in keep])
    return buffer.getvalue()
